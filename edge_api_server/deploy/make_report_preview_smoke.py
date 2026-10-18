"""Rebuild the small, non-confidential PPTX installation test fixture."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

target = Path(__file__).with_name("report_preview_smoke.pptx")
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
for n, (text, font) in enumerate([
    ("RIST TEM", "Arial"), ("한글 미리보기 검증", "Noto Sans CJK KR"),
    ("시료명: 검증시료", "NanumGothic"), ("두께: 2.50 nm", "Malgun Gothic"),
]):
    box = slide.shapes.add_textbox(Inches(1), Inches(.7+n*1.1), Inches(8), Inches(.8))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(28)
deck.save(target)
