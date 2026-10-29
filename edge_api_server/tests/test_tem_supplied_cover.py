"""The supplied TEM cover changes artwork only, including after revisions."""
from copy import deepcopy
import hashlib
import posixpath
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from pptx import Presentation

from ahn import ppt_report
from ahn.report_branding import NS, brand_assets, brand_pictures
from ahn.report_cover import COVER_NAME, SUPPLIED_COVER_NAME, shape_xml, supplied_cover_asset
from ahn.report_traceability import TRACE_NAME, cover_shapes, traceability_shapes
from ahn.tests.test_ppt_report import _project_with_sections
from app.report import canvas, tem_canvas, tem_editable


METADATA = dict(title="TEM 검토 보고서", client="RIST 의뢰기관", requester="홍길동",
                analyst="분석 담당자", completionDate="2026-10-29", requestNumber="REQ-TEM",
                limsExperimentCode="TEM-TEST", equipmentCode="EQ-TEM", operatorId="OP-TEM")


@pytest.fixture
def report(tmp_path):
    project = _project_with_sections(tmp_path, {"tem"})
    project["report_metadata"] = dict(METADATA)
    return ppt_report.build_pptx(project, tmp_path / "source.pptx")


def text_shapes(root):
    return {s.find("p:nvSpPr/p:cNvPr", NS).get("name"): s
            for s in root.findall("p:cSld/p:spTree/p:sp", NS)}


def picture_bytes(archive, page, picture):
    rid = picture.find("p:blipFill/a:blip", NS).get("{" + NS["r"] + "}embed")
    rels = ET.fromstring(archive.read(tem_canvas._rels_name(page)))
    target = next(r.get("Target") for r in rels if r.get("Id") == rid)
    return archive.read(posixpath.normpath(posixpath.join(posixpath.dirname(page), target)))


@pytest.mark.parametrize("use_template", [True, False])
def test_only_cover_artwork_changes_and_original_logo_remains(tmp_path, monkeypatch, use_template):
    if not use_template:
        monkeypatch.setattr(ppt_report, "TEMPLATE_PATH", tmp_path / "missing-template.pptx")
    project = _project_with_sections(tmp_path, {"tem"})
    project["report_metadata"] = dict(METADATA)
    path = ppt_report.build_pptx(project, tmp_path / "report.pptx")
    prs = Presentation(path)
    width, height = prs.slide_width, prs.slide_height
    values = dict(METADATA, sampleNames="TEM-A", experimentCode="TEM-TEST",
                  sourceSummary="analysis-result.json · 원자료 경로 및 이미지·EDS 대응 목록",
                  revision="최초 생성본")
    with ZipFile(path) as archive:
        pages = tem_editable._slides(archive)
        assert len(pages) == 3
        cover, trace, body = [ET.fromstring(archive.read(name)) for name in pages]
        assert cover.find("p:cSld", NS).get("name") == COVER_NAME
        assert trace.find("p:cSld", NS).get("name") == TRACE_NAME
        assert all(root.get("showMasterSp") is None for root in (cover, trace, body))
        for root, layout in ((cover, cover_shapes), (trace, traceability_shapes)):
            expected = {s["name"]: ET.fromstring(shape_xml(s)) for s in
                        layout("TEM", METADATA["title"], values, width, height)}
            actual = text_shapes(root)
            assert actual.keys() == expected.keys()
            for name in expected:
                assert ET.tostring(actual[name]) == ET.tostring(expected[name]), name
        pictures = cover.findall("p:cSld/p:spTree/p:pic", NS)
        assert [p.find("p:nvPicPr/p:cNvPr", NS).attrib for p in pictures] == [
            dict(id="1000", name=SUPPLIED_COVER_NAME), dict(id="1001", name="Template RIST logo")]
        artwork, logo = pictures
        assert picture_bytes(archive, pages[0], artwork) == supplied_cover_asset()
        xfrm = artwork.find("p:spPr/a:xfrm", NS)
        assert xfrm.find("a:off", NS).attrib == dict(x="0", y="0")
        assert xfrm.find("a:ext", NS).attrib == dict(cx=str(width), cy=str(height))
        original_logo = next(p for p in brand_pictures(width, height) if p.name == "RIST logo")
        assert picture_bytes(archive, pages[0], logo) == brand_assets()[original_logo.media]
        expected_logo = ET.fromstring(original_logo.xml(1001, "unused"))
        assert ET.tostring(logo.find("p:spPr", NS)).strip() == ET.tostring(expected_logo.find("p:spPr", NS)).strip()
        visible_shapes = list(cover.find("p:cSld/p:spTree", NS))[2:]
        assert visible_shapes[:2] == pictures  # Both pictures are behind the original text.
        ids = [node.get("id") for node in cover.findall(".//p:cNvPr", NS)]
        assert len(ids) == len(set(ids))
        assert not any(node.get("name") == SUPPLIED_COVER_NAME
                       for root in (trace, body) for node in root.findall(".//p:cNvPr", NS))


@pytest.mark.parametrize("patch", [dict(hidden=True), dict(x=.2), dict(asset="a" * 64)])
def test_cover_artwork_cannot_be_deleted_moved_or_replaced(report, tmp_path, patch):
    with ZipFile(report) as archive:
        payload = tem_canvas._read(archive)[0]
    assert not any(item["name"] in {SUPPLIED_COVER_NAME, "Template RIST logo"}
                   for _, item in canvas.objects(payload)[1])
    oid = "shape-" + hashlib.sha256((SUPPLIED_COVER_NAME + ":0").encode()).hexdigest()[:20]
    document = SimpleNamespace(page_edits={}, canvas_asset_root=tmp_path,
                               canvas_edits={payload["pageId"]: [dict(id=oid, **patch)]})
    with pytest.raises(ValueError, match="편집할 개체가 없습니다"):
        canvas.apply(document, [payload], validate_only=True)


def test_revisions_preserve_background_logo_body_and_original_edit_bindings(report, tmp_path):
    original_model = tem_editable.initial_model(report, {})
    _, original_layout, original_bindings = tem_editable.inspect_pptx(report)
    assert original_model["completionDate"] == "2026-10-29"
    with ZipFile(report) as archive:
        original_pages = tem_editable._slides(archive)
        original_body = ET.tostring(ET.fromstring(archive.read(original_pages[2])))
        original_art = [ET.tostring(p) for p in ET.fromstring(archive.read(original_pages[0]))
                        .findall("p:cSld/p:spTree/p:pic", NS)]
        immutable = {name: archive.read(name) for name in archive.namelist()
                     if name.startswith(("ppt/media/", "ppt/slideLayouts/", "ppt/slideMasters/"))}
    for index in range(2):
        model = tem_editable.initial_model(report, {})
        model.update(title=f"검토 보고서 {index}", completionDate=f"2026-11-0{index + 1}")
        client = next(s for s in model["sections"] if s["heading"] == "표지 · 의뢰기관")
        client["paragraphs"] = [f"기관 수정 {index}" if index == 0 else ""]
        target = tmp_path / f"revision-{index}.pptx"
        tem_editable.revise_pptx(report, target, model)
        tem_editable.stamp_record(target, f"revision-{index}", "original-job")
        reopened = tem_editable.initial_model(target, {})
        _, layout, bindings = tem_editable.inspect_pptx(target)
        assert reopened["title"] == model["title"]
        assert reopened["completionDate"] == model["completionDate"]
        assert bindings == original_bindings
        assert [r["refs"] for r in layout["pages"][0]["regions"]] == [
            r["refs"] for r in original_layout["pages"][0]["regions"]]
        with ZipFile(target) as archive:
            pages = tem_editable._slides(archive)
            assert pages == original_pages
            cover = ET.fromstring(archive.read(pages[0]))
            assert [ET.tostring(p) for p in cover.findall("p:cSld/p:spTree/p:pic", NS)] == original_art
            assert ET.tostring(ET.fromstring(archive.read(pages[2]))) == original_body
            assert all(archive.read(name) == content for name, content in immutable.items())
            shapes = text_shapes(cover)
            assert tem_editable._text(shapes["Cover completionDate"].find("p:txBody", NS)) == "완료일: " + model["completionDate"]
            assert tem_editable._text(shapes["Cover client"].find("p:txBody", NS)) == "의뢰기관: " + (client["paragraphs"][0] or "미기재")
        report = target


def test_legacy_cover_stays_without_supplied_artwork(report, tmp_path):
    prs = Presentation(report)
    cover = prs.slides[0]
    for shape in list(cover.shapes):
        if shape.name in {SUPPLIED_COVER_NAME, "Template RIST logo"}:
            ppt_report._remove_shape(shape)
    legacy = tmp_path / "legacy.pptx"
    prs.save(legacy)
    model = tem_editable.initial_model(legacy, {})
    target = tmp_path / "legacy-revised.pptx"
    tem_editable.revise_pptx(legacy, target, deepcopy(model))
    assert tem_editable.initial_model(target, {}) == model
    reopened = Presentation(target)
    assert len(reopened.slides) == len(prs.slides)
    assert reopened.slides[0].element.get("showMasterSp") is None
    assert not any(shape.name == SUPPLIED_COVER_NAME for shape in reopened.slides[0].shapes)
