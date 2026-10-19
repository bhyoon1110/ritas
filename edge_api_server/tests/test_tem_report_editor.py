from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta
from io import BytesIO
import json
from pathlib import Path
import time
from types import SimpleNamespace
from zipfile import ZipFile

from fastapi.testclient import TestClient
import pytest

from app import ahn_web, tem_report_editor
from app import preview_report
from app.auth import AuthContext
from app.report import tem_editable
from app.report_queue import ReportQueueError
from app.storage import atomic_write_json
from ahn.ppt_report import build_pptx
from ahn.tests.test_ppt_report import _project_with_sections


def editable(value):
    value = deepcopy(value)
    for section in value["sections"]:
        section.pop("columns", None)
        section.pop("source", None)
    return value


@pytest.fixture
def report(tmp_path, monkeypatch, tem_unit_runtime):
    app = ahn_web.create_tem_preview_app()
    app.state.actor = None
    @app.middleware("http")
    async def identity(request, call_next):
        request.state.auth_context = app.state.actor
        return await call_next(request)
    root = ahn_web.tem_history.new_workspace(app.state.settings)
    inputs = root / "input"
    inputs.mkdir(parents=True)
    job = ahn_web._create_ahn_job(inputs, root, None, app.state.settings, owner_user_id="author")
    job.output_dir.mkdir()
    project = _project_with_sections(inputs, {"tem", "coating"})
    build_pptx(project, job.pptx_path)
    atomic_write_json(job.analysis_path, project)
    manifest = {"summary": {"temSampleCount": 1, "coatingSampleCount": 1}, "pptx": str(job.pptx_path)}
    atomic_write_json(job.manifest_path, manifest)
    tem_report_editor.prepare_review(job, manifest)
    ahn_web._build_package(job.output_dir, job.package_path)
    ahn_web._set_job_state(job, status="completed", manifest=manifest)
    with TestClient(app) as client:
        yield SimpleNamespace(app=app, client=client, job=job, url="/api/v1/tem/report/jobs/" + job.job_id)
    with ahn_web._ahn_report_jobs_lock:
        for key, item in list(ahn_web._ahn_report_jobs.items()):
            if item is job or item.owner_user_id == "author":
                ahn_web._ahn_report_jobs.pop(key)


def wait(h, payload):
    for _ in range(500):
        result = h.client.get("/api/v1/tem/report/jobs/" + payload["jobId"])
        assert result.status_code == 200
        result = result.json()
        if result["status"] in {"completed", "failed"}:
            return result
        time.sleep(.01)
    pytest.fail("revision did not finish")


def test_completed_history_reopens_editor_and_slides_after_registry_reload(report):
    h = report
    history = ahn_web.tem_history
    before = h.client.get(h.url + "/editor-data").json()["document"]
    with ahn_web._ahn_report_jobs_lock:
        ahn_web._ahn_report_jobs.pop(h.job.job_id)
        history._loaded.discard(history.root(h.app.state.settings))
    response = h.client.get(h.url + "/editor-data")
    assert response.status_code == 200
    assert response.json()["document"] == before
    assert h.client.get(h.url + "/edit-layout").status_code == 200
    assert h.client.get(h.url + "/edit-slides/0").headers["content-type"] == "image/jpeg"
    assert h.client.get(h.url + "/download").status_code == 200
    revised = h.client.post(h.url + "/revisions", json=editable(before),
                            headers={"X-Requested-With": "RIST-Report-Editor"})
    assert revised.status_code == 202
    assert wait(h, revised.json())["status"] == "completed"


def test_tem_revision_edits_visible_text_table_email_and_preserves_original(report, monkeypatch):
    h = report
    before = h.job.package_path.read_bytes()
    page = h.client.get(h.url + "/edit")
    assert page.status_code == 200 and "frame-ancestors 'self'" in page.headers["content-security-policy"]
    data = h.client.get(h.url + "/editor-data").json()["document"]
    assert "TEM 1건" in data["auxiliary"]["email_body"]
    layout = h.client.get(h.url + "/edit-layout").json()
    assert len(layout["pages"]) == 3  # New cover + the two original body pages.
    assert h.client.get(h.url + "/edit-slides/0").headers["content-type"] == "image/jpeg"
    assert h.client.get(h.url + "/edit-slides/-1").status_code == 404
    assert h.client.get(h.url + "/edit-slides/999").status_code == 404
    changes = editable(data)
    changes["transfer"] = dict(requestNumber="REQ-NEW", limsExperimentCode="LIMS-TEM", equipmentCode="EQ-1", operatorId="OP-1")
    changes["auxiliary"] = {"email_subject": "{{의뢰번호}} 분석 결과", "email_body": "검토 완료한 TEM 보고서입니다."}
    text = next(s for s in changes["sections"] if s["paragraphs"] and "TEM-A" in s["paragraphs"][0])
    text["paragraphs"][0] = "교정한 시료명 <literal>"
    table = next(s for s in changes["sections"] if s["rows"])
    table["rows"][1][-1] = "2.34"
    table["rows"][-1][-1] = "2.67"
    def no_analysis(**kwargs):
        pytest.fail("Review must not rerun analysis or OCR")
    monkeypatch.setattr(ahn_web, "build_outputs", no_analysis)
    response = h.client.post(h.url + "/revisions", json=changes, headers={"X-Requested-With": "RIST-Report-Editor"})
    assert response.status_code == 202, response.text
    result = wait(h, response.json())
    assert result["status"] == "completed", result
    assert result["parentJobId"] == h.job.job_id
    base = "/api/v1/tem/report/jobs/" + result["jobId"]
    reopened = h.client.get(base + "/editor-data").json()["document"]
    assert next(s for s in reopened["sections"] if s["sectionId"] == table["sectionId"])["columns"]
    assert reopened["auxiliary"]["email_subject"] == "{{의뢰번호}} 분석 결과"
    package = h.client.get(base + "/download").content
    with ZipFile(BytesIO(package)) as archive, ZipFile(h.job.pptx_path) as original:
        assert "review/editor-state.json" not in archive.namelist()
        assert "REQ-NEW 분석 결과" in archive.read("email_body.md").decode()
        assert json.loads(archive.read("analysis-result.json")) == json.loads(h.job.analysis_path.read_text())
        assert json.loads(archive.read("report-review.json"))["transfer"] == changes["transfer"]
        with ZipFile(BytesIO(archive.read("tem-report.pptx"))) as revised:
            cover = revised.read("ppt/slides/slide1.xml").decode()
            for value in changes["transfer"].values():
                assert value in cover
            for name in original.namelist():
                if name.startswith(("ppt/media/", "ppt/slideLayouts/", "ppt/slideMasters/")):
                    assert revised.read(name) == original.read(name), name
            texts = b"".join(revised.read(n) for n in revised.namelist() if n.startswith("ppt/slides/") and n.endswith(".xml")).decode()
            assert "교정한 시료명 &lt;literal&gt;" in texts and "2.34" in texts
            assert any(n.startswith("ppt/media/tem-review-") for n in revised.namelist())
    assert h.job.package_path.read_bytes() == before
    assert h.job.active_revisions == 0
    captured = {}
    def queue(**kwargs):
        captured.update(kwargs)
        return {"queued": True}
    monkeypatch.setattr(ahn_web, "send_preview_report_package", queue)
    payload = dict(requestNumber="REQ-NEW", experimentCode="LIMS-TEM", equipmentCode="EQ-1", operatorId="OP-1")
    assert h.client.post(base + "/send", json=payload).status_code == 200
    assert captured["job"].job_id == result["jobId"]
    assert captured["job"].package_path.read_bytes() == package


@pytest.mark.parametrize("alter", ["shape", "table", "paragraph", "variable", "mail", "date"])
def test_tem_rejects_forged_structure_and_invalid_values(report, alter):
    h = report
    value = editable(h.client.get(h.url + "/editor-data").json()["document"])
    if alter == "shape": value["sections"][0]["sectionId"] = "../../etc/password"
    if alter == "table": next(s for s in value["sections"] if s["rows"])["rows"].append(["new"])
    if alter == "paragraph": next(s for s in value["sections"] if s["paragraphs"])["paragraphs"].append("extra")
    if alter == "variable": value["title"] = "{{unknown}}"
    if alter == "mail": value["auxiliary"]["email_body"] = "x" * 4001
    if alter == "date": value["completionDate"] = "2026-99-99"
    result = h.client.post(h.url + "/revisions", json=value, headers={"X-Requested-With": "RIST-Report-Editor"})
    assert result.status_code == 422, result.text


def test_tem_csrf_size_limit_and_failure_keep_original(report, monkeypatch):
    h = report
    value = editable(h.client.get(h.url + "/editor-data").json()["document"])
    assert h.client.post(h.url + "/revisions", json=value).status_code == 403
    assert h.client.post(h.url + "/revisions", json=value, headers={"X-Requested-With": "RIST-Report-Editor", "Origin": "https://other.invalid"}).status_code == 403
    assert h.client.post(h.url + "/revisions", content=b"x" * (2*1024*1024+1), headers={"X-Requested-With": "RIST-Report-Editor"}).status_code == 413
    original = h.job.package_path.read_bytes()
    monkeypatch.setattr(tem_report_editor, "convert_pptx_to_pdf", lambda *a: (_ for _ in ()).throw(RuntimeError("test conversion failure")))
    result = h.client.post(h.url + "/revisions", json=value, headers={"X-Requested-With": "RIST-Report-Editor"})
    assert wait(h, result.json())["status"] == "failed"
    assert h.job.package_path.read_bytes() == original and h.job.active_revisions == 0


def test_tem_author_admin_and_project_access(report):
    h = report
    h.app.state.settings = replace(h.app.state.settings, auth_enabled=True)
    actor = AuthContext(user_id="author", login_id="author", email=None, display_name="Author", status="ACTIVE",
        session_id="test", session_expires_at=datetime.now()+timedelta(hours=1), sso_authenticated_at=None,
        projects=frozenset({"TEM"}), roles=frozenset(), sso_identity=None)
    for user, projects, roles, status in [("author", {"TEM"}, set(), 200), ("other", {"TEM"}, set(), 403),
            ("author", {"FTIR"}, set(), 403), ("admin", set(), {"ADMIN"}, 200)]:
        h.app.state.actor = replace(actor, user_id=user, projects=frozenset(projects), roles=frozenset(roles))
        for suffix in ("", "/edit", "/editor-data", "/edit-layout", "/edit-slides/0", "/download", "/download/pptx", "/preview"):
            assert h.client.get(h.url+suffix).status_code == status, (user, suffix)
    h.app.state.actor = actor
    payload = dict(requestNumber="REQ", experimentCode="TEM", equipmentCode="EQ", operatorId="OP")
    assert h.client.post(h.url + "/send", json=payload).status_code == 403  # SSO still required.


def test_tem_missing_preview_cannot_send_and_does_not_destroy_ppt(report):
    h = report
    h.job.review_error = "Preview unavailable"
    assert h.client.get(h.url + "/edit-layout").status_code == 503
    assert h.client.get(h.url + "/download").status_code == 200
    assert h.client.post(h.url + "/send", json=dict(requestNumber="REQ", experimentCode="TEM", equipmentCode="EQ", operatorId="OP")).status_code == 409


def test_tem_queue_uses_saved_metadata_and_never_enqueues_a_stale_identity(report, monkeypatch):
    h = report
    metadata = dict(requestNumber="REQ", limsExperimentCode="METHOD", equipmentCode="EQ", operatorId="OP")
    h.job.transfer_metadata = metadata
    h.app.state.database = object()
    queued = []
    monkeypatch.setattr(preview_report, "persist_preview_report_package", lambda **kw: kw["source_path"])
    monkeypatch.setattr(preview_report, "enqueue_report_package", lambda **kw: queued.append(kw) or {"queued": True})
    payload = dict(requestNumber="REQ", experimentCode="METHOD", equipmentCode="EQ", operatorId="STALE")
    assert h.client.post(h.url + "/send", json=payload).status_code == 409
    assert not queued
    payload["operatorId"] = "OP"
    result = h.client.post(h.url + "/send", json=payload)
    assert result.status_code == 200 and result.json()["sent"] is False
    assert queued[0]["package_path"] == h.job.package_path
    assert queued[0]["report_id"] == h.job.job_id


def test_tem_failed_setup_releases_pending_slot_and_source_pin(report, monkeypatch):
    from threading import BoundedSemaphore
    h = report
    slots = BoundedSemaphore(1)
    monkeypatch.setattr(ahn_web, "_ahn_pending_slots", slots)
    monkeypatch.setattr(ahn_web.tem_history, "new_workspace", lambda *args: (_ for _ in ()).throw(OSError("test disk full")))
    value = editable(h.client.get(h.url + "/editor-data").json()["document"])
    with pytest.raises(OSError, match="test disk full"):
        h.client.post(h.url + "/revisions", json=value, headers={"X-Requested-With": "RIST-Report-Editor"})
    assert slots.acquire(blocking=False)
    slots.release()
    assert h.job.active_revisions == 0 and h.job.pptx_path.is_file()


def test_tem_busy_revision_queue_returns_retryable_status_without_new_job(report, monkeypatch):
    from threading import BoundedSemaphore
    h = report
    monkeypatch.setattr(ahn_web, "_ahn_pending_slots", BoundedSemaphore(0))
    value = editable(h.client.get(h.url + "/editor-data").json()["document"])
    assert h.client.post(h.url + "/revisions", json=value, headers={"X-Requested-With": "RIST-Report-Editor"}).status_code == 429
    assert h.job.active_revisions == 0


def test_tem_repeated_coating_edits_preserve_geometry_and_do_not_duplicate_zip_members(report, tmp_path):
    h = report
    model = tem_report_editor._state(h.job)
    rows = next(s["rows"] for s in model["sections"] if s["rows"])
    rows[1][1] = "2.34"
    first, second = tmp_path / "first.pptx", tmp_path / "second.pptx"
    tem_editable.revise_pptx(h.job.pptx_path, first, model)
    rows[1][1] = "2.50"
    tem_editable.revise_pptx(first, second, model)
    original_layout = tem_editable.inspect_pptx(h.job.pptx_path)[1]
    sections, layout, _ = tem_editable.inspect_pptx(second)
    assert layout == original_layout
    assert next(s["rows"] for s in sections if s["rows"])[1][1] == "2.50"
    with ZipFile(second) as z:
        assert z.testzip() is None
        assert len(z.namelist()) == len(set(z.namelist()))


def test_tem_native_table_and_multiline_text_are_editable_without_moving_shapes(report, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(report.job.pptx_path)
    table = prs.slides[0].shapes.add_table(2, 2, Inches(6), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Element"
    table.cell(0, 1).text = "Weight %"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "40.0"
    text = prs.slides[0].shapes.add_textbox(Inches(6), Inches(4), Inches(3), Inches(1))
    text.text = "first\vsecond"  # DrawingML line break, not another paragraph.
    source, target = tmp_path / "native.pptx", tmp_path / "native-edited.pptx"
    prs.save(source)
    model = tem_editable.initial_model(source, {})
    table_section = next(s for s in model["sections"] if s["rows"] and s["rows"][0][0] == "Element")
    table_section["rows"][1][1] = "42.5"
    text_section = next(s for s in model["sections"] if s["paragraphs"] == ["first\nsecond"])
    text_section["paragraphs"][0] = "교정\n설명"
    before = tem_editable.inspect_pptx(source)[1]
    tem_editable.revise_pptx(source, target, model)
    sections, layout, _ = tem_editable.inspect_pptx(target)
    assert layout == before
    assert next(s for s in sections if s["sectionId"] == table_section["sectionId"])["rows"][1][1] == "42.5"
    assert next(s for s in sections if s["sectionId"] == text_section["sectionId"])["paragraphs"] == ["교정\n설명"]
