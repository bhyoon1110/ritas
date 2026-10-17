"""16:9 landscape page model: one geometry, header and footer as page items, rect items."""

import io

from .figure_fixtures import make_document, make_project, make_sample
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pypdf import PdfReader

from app.xrd_workbench.render import theme
from app.xrd_workbench.render.html import render_html
from app.xrd_workbench.render.layout import BOTTOM, LEFT, WIDTH, build_pages
from app.xrd_workbench.render.pdf import render_pdf
from app.xrd_workbench.render.pptx import render_pptx


def pages_for(n=2, title=None, report_number=None):
    samples = [make_sample(f's{i}', f'시료 {i}', [], []) for i in range(n)]
    project = make_project(samples)
    document = make_document(project, include=False)
    if title:
        document['metadata']['title'] = title
    if report_number:
        document['metadata']['report_number'] = report_number
    return build_pages(project, document, lambda fid: b''), document


def body(page):
    return [i for i in page['items'] if not i.get('role')]


def test_the_page_is_16_by_9_widescreen_with_symmetric_margins():
    assert (round(theme.PAGE_W, 3), theme.PAGE_H) == (338.667, 190.5)
    assert abs(2 * LEFT + WIDTH - theme.PAGE_W) < 0.01


def test_every_page_starts_with_a_full_width_title_band_and_white_title():
    pages, _ = pages_for()
    assert pages[0]['kind'] == 'cover'  # the cover has no title band
    for page in pages[1:]:
        band = next(i for i in page['items'] if i['type'] == 'rect')
        assert (band['x'], band['y'], band['w']) == (0, 0, theme.PAGE_W)
        assert band['fill'] == theme.NAVY and band['role'] == 'header'
        title = next(i for i in page['items'] if i.get('role') == 'header' and i['type'] == 'text')
        assert title['color'] == theme.WHITE and title['text'].replace('\n', '') == page['title']
        assert page['items'][0] is band


def test_a_long_title_wraps_and_the_band_and_content_grow_with_it():
    pages, _ = pages_for(title='제목')
    long_name = 'WWWWWWWWWW ' * 17
    project = make_project([make_sample('s0', long_name, [], [])])
    document = make_document(project, include=False)
    document['samples'][0]['display_name'] = long_name
    document['samples'][0]['manual_note'] = '본문 문장'
    sample_page = next(
        p for p in build_pages(project, document, lambda fid: b'') if p['kind'] == 'sample'
    )
    band = next(i for i in sample_page['items'] if i['type'] == 'rect')
    assert band['h'] > theme.HEADER_MIN_H
    assert min(i['y'] for i in body(sample_page)) >= band['h'] + 6 - 0.01
    assert pages[0]['content_top'] == theme.TOP


def test_footer_has_the_report_number_and_a_right_aligned_page_number():
    pages, _ = pages_for(report_number='RN-001')
    for number, page in enumerate(pages, 1):
        footer = [i for i in page['items'] if i.get('role') == 'footer']
        assert any(i['text'] == 'RN-001' and i['x'] == LEFT for i in footer)
        count = next(i for i in footer if i['text'] == f'{number} / {len(pages)}')
        assert count['align'] == 'right'
        assert count['x'] + count['w'] <= theme.PAGE_W - LEFT + 0.01


def test_body_items_stay_inside_the_page_and_above_the_footer():
    pages, _ = pages_for(5)
    for page in pages:
        for item in body(page):
            assert 0 <= item['x'] and item['x'] + item['w'] <= theme.PAGE_W + 0.01
            assert item['y'] + item['h'] <= BOTTOM + 0.01


def test_pdf_pages_are_16_by_9_at_the_theme_size():
    pages, document = pages_for()
    reader = PdfReader(io.BytesIO(render_pdf(pages, document)))
    box = reader.pages[0].mediabox
    assert abs(float(box.width) - theme.PAGE_W / 25.4 * 72) < 0.5
    assert abs(float(box.height) - theme.PAGE_H / 25.4 * 72) < 0.5


def test_html_pages_use_the_theme_size_and_draw_the_band_as_an_item():
    pages, document = pages_for()
    html = render_html(pages, document, 'package')
    assert f'width:{theme.PAGE_W}mm' in html and f'height:{theme.PAGE_H}mm' in html
    assert theme.NAVY in html and 'pageHeader' not in html and 'pageFooter' not in html
    assert 'text-align:right' in html


def test_pptx_slides_are_16_by_9_with_a_band_shape_and_right_aligned_page_number():
    pages, document = pages_for()
    deck = Presentation(io.BytesIO(render_pptx(pages, document)))
    assert (deck.slide_width, deck.slide_height) == (12192000, 6858000)
    shapes = list(deck.slides[1].shapes)  # slide 1 is the cover, which has no band
    bands = [
        s
        for s in shapes
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.width == deck.slide_width
    ]
    assert bands and str(bands[0].fill.fore_color.rgb) == theme.NAVY.lstrip('#').upper()
    numbers = [s for s in shapes if s.has_text_frame and s.text_frame.text == f'2 / {len(pages)}']
    assert numbers and numbers[0].text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


def test_renderers_need_no_header_or_footer_fields_on_the_page():
    page = {'kind': 'x', 'title': 't', 'items': [], 'number': 1, 'total': 1}
    document = {'metadata': {'title': 't', 'author': ''}}
    assert PdfReader(io.BytesIO(render_pdf([page], document))).pages
    assert 'class="page"' in render_html([page], document, 'package')
    assert len(Presentation(io.BytesIO(render_pptx([page], document))).slides) == 1


def test_the_cover_lists_summarised_samples_in_two_columns_row_by_row():
    pages, _ = pages_for(4)
    cover = next(p for p in pages if p['kind'] == 'summary')
    headings = [
        i
        for i in cover['items']
        if i['type'] == 'text'
        and not i.get('role')
        and i['text'] in ('시료 0', '시료 1', '시료 2', '시료 3')
    ]
    assert [h['text'] for h in headings] == ['시료 0', '시료 1', '시료 2', '시료 3']
    column_w = (WIDTH - theme.COVER_GAP) / 2
    assert [h['x'] for h in headings] == [LEFT, LEFT + column_w + theme.COVER_GAP] * 2
    assert headings[0]['y'] == headings[1]['y'] and headings[2]['y'] == headings[3]['y']
    assert headings[2]['y'] > headings[0]['y']
    assert all(h['w'] == column_w for h in headings)


def test_a_cover_summary_that_does_not_fit_is_still_refused_with_the_same_message():
    import pytest

    samples = [
        make_sample(f's{i}', f'시료 {i}', [(f'P{j}', 'c', 'D') for j in range(5)], [])
        for i in range(4)
    ]
    project = make_project(samples)
    document = make_document(project, include=False)
    document['summary'] = ('매우 긴 요약 문장입니다. ' * 100)[:1499]
    for entry in document['samples']:
        entry['display_name'] = 'W' * 200
    with pytest.raises(
        ValueError, match='첫 장 요약이 페이지를 넘습니다|첫 장 내용이 페이지를 넘습니다'
    ):
        build_pages(project, document, lambda fid: b'')


def test_printed_html_keeps_the_title_band_background():
    pages, document = pages_for()
    html = render_html(pages, document, 'package')
    assert 'print-color-adjust:exact' in html
