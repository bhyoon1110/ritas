"""The supplied three-layout deck is a presentation template, never example analysis data."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import io
import posixpath
from zipfile import ZipFile

from lxml import etree as E
from openpyxl import load_workbook

from app.xrd_workbench.core.report import new_document
from app.xrd_workbench.core.phase_class import phase_ref
from app.xrd_workbench.render.posco_pptx import (
    TEMPLATE, NS, Deck, render_report_pptx, named,
    SUMMARY_TITLE, QUANTITY_TITLE, PHASE_TITLE,
)
from .figure_fixtures import make_project
from .test_sample_slide import sample, IMAGE


def fixture():
    project=make_project([sample('A',quantities=0),sample('B',quantities=0),sample('Mix',quantities=3)])
    for quantity in project['samples'][2]['quantities']:
        quantity['value']=float(quantity['raw'].split('(')[0])
    document=new_document(project)
    document['figure']['include']=False
    document['quantity_table']['include']=False
    document['metadata'].update(title='POSCO 합성 검증',report_number='REQ-123',requester='시험 의뢰자',
                                equipment_code='XRD-01',operator_id='QA-OP')
    document['samples'][0]['manual_note']='측정값을 그대로 보존하는 검토 문장'
    for cls,phase in zip(('major','trace','presumed'),project['samples'][0]['phases']):
        document['phase_classes']['labels'].append(dict(sample_id='A',phase=phase_ref(phase),
                                                        origin='manual',**{'class':cls}))
    return project,document


def deck(data):
    z=ZipFile(io.BytesIO(data))
    names=sorted((n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')),
                 key=lambda n:int(n.split('slide')[-1].split('.')[0]))
    slides=[E.fromstring(z.read(n)) for n in names]
    text='\n'.join(''.join(s.itertext()) for s in slides)
    return z,slides,text


def test_reference_background_dimensions_and_sanitized_example_data():
    project,document=fixture()
    before=deepcopy((project,document))
    z,slides,text=deck(render_report_pptx(project,document,lambda _:IMAGE))
    with ZipFile(TEMPLATE) as original:
        for name in original.namelist():
            if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/','ppt/media/')):
                assert z.read(name)==original.read(name),name
    size=E.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz',NS)
    assert (size.get('cx'),size.get('cy'))==('9906000','6858000')
    assert 'Zinc Manganese' not in text and '00-062-0414' not in text and '{Sample' not in text
    assert len(slides)==5
    assert [''.join(named(s,'Title')[0].xpath('.//a:t/text()',namespaces=NS)) for s in slides]==[
        'XRD 분석 보고서', '분석 이력 요약', SUMMARY_TITLE, QUANTITY_TITLE, PHASE_TITLE,
    ]
    assert '검토 문장' not in text  # prose/LLM is preserved in the document, not the rules-only summary
    assert all(value not in text for value in ('보고서 정보','정량 수치','결정상 분류 기준','긴 제목 전체'))
    assert before==(project,document)
    assert len(slides[4].findall('.//a:tbl',NS))==2  # phase-only samples share the reference layout
    assert len(slides[3].findall('.//a:tbl',NS))==2  # WPPF includes a native quantity table
    assert '[주요상] Phase0' in text and '[미량상] Phase1' in text and '[추정상] Phase2' in text
    assert '[미분류] Phase0' in text  # no implicit scientific classification


def test_native_chart_workbook_and_cache_preserve_quantities_and_raw_precision():
    project,document=fixture()
    z,_,text=deck(render_report_pptx(project,document,lambda _:IMAGE))
    chart=E.fromstring(z.read('ppt/charts/chart1.xml'))
    values=chart.findall('.//c:numCache/c:pt/c:v',NS)
    assert [float(v.text) for v in values]==[10,11.1,12.2]
    book=load_workbook(io.BytesIO(z.read('ppt/embeddings/quantity1.xlsx')),data_only=True)
    assert list(book.active.values)[1:]==[('Phase0',10,'10.0(1)'),('Phase1',11.1,'11.1(2)'),('Phase2',12.2,'12.2(3)')]
    assert all(v in text for v in ('10.0(1) wt%','11.1(2) wt%','12.2(3) wt%'))
    assert '합계 33.3 wt%' in text
    # Every internal relationship resolves, including charts and workbooks.
    for name in z.namelist():
        if not name.endswith('.rels'):continue
        root=E.fromstring(z.read(name))
        parent=posixpath.dirname(name).replace('/_rels','') if name!='_rels/.rels' else ''
        for rel in root:
            if rel.get('TargetMode')=='External':continue
            target=posixpath.normpath(posixpath.join(parent,rel.get('Target')))
            assert target in z.namelist(),(name,target)


def test_more_samples_keep_only_reference_layouts_and_do_not_mutate_full_notes():
    project,document=fixture()
    project['samples']= [sample(str(i),n_phases=8,quantities=0) for i in range(5)]
    document=new_document(project)
    document['figure']['include']=False;document['quantity_table']['include']=False
    for i,e in enumerate(document['samples']):e['summary_included']=i<4
    document['samples'][0]['manual_note']='QA-START '+('긴 검토문장입니다. '*180)+' QA-END'
    before=deepcopy((project,document))
    z,slides,text=deck(render_report_pptx(project,document,lambda _:IMAGE,include_sources=False))
    assert before==(project,document)
    assert 'QA-START' not in text and '뒤쪽' not in text
    assert document['samples'][0]['manual_note'].endswith('QA-END')
    assert all('시료 '+str(i) in text for i in range(5))
    assert 'Phase7' in text and '상 목록 1–3' not in text
    titles=[''.join(named(s,'Title')[0].xpath('.//a:t/text()',namespaces=NS)) for s in slides]
    assert titles.count(PHASE_TITLE) == 3  # five whole sample tables, paired per slide
    assert PHASE_TITLE+' (계속)' not in titles
    assert all(t.removesuffix(' (계속)') in ('XRD 분석 보고서','분석 이력 요약',SUMMARY_TITLE,PHASE_TITLE) for t in titles)
    for n in z.namelist():
        if n.endswith('.rels'):assert b'TargetMode="External"' not in z.read(n)
    for slide in slides:
        for shape in slide.findall('.//p:sp',NS):
            xf=shape.find('p:spPr/a:xfrm',NS)
            if xf is not None:assert int(xf[0].get('y'))+int(xf[1].get('cy'))<=6858000


def test_concurrent_templates_do_not_leak_titles_or_data():
    project,document=fixture()
    def output(title):
        value=deepcopy(document);value['metadata']['title']=title
        z,_,_=deck(render_report_pptx(project,value,lambda _:IMAGE))
        return z.read('docProps/core.xml').decode()
    titles=['동시 검증 AAA','동시 검증 BBB']
    with ThreadPoolExecutor(2) as pool:values=list(pool.map(output,titles))
    for i,value in enumerate(values):assert titles[i] in value and titles[1-i] not in value


def test_source_artwork_is_exact_and_user_quantity_exclusions_remain_missing():
    project,document=fixture()
    from .test_sample_slide import source
    project['samples'][2]['sources'].append(source('Mix','quantity_image','quant.png',IMAGE))
    document['samples'][2]['include_quantity_image']=True
    z,_,_=deck(render_report_pptx(project,document,lambda _:IMAGE))
    assert IMAGE in [z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
    assert not any(n.startswith('ppt/charts/chart') for n in z.namelist())


def test_missing_inconsistent_or_mixed_unit_quantities_are_not_silently_graphed():
    for change in ({'value':None},{'value':-1},{'value':999},{'unit':'vol%'}):
        project,document=fixture()
        project['samples'][2]['quantities'][0].update(change)
        document=new_document(project)
        document['figure']['include']=False;document['quantity_table']['include']=False
        z,_,text=deck(render_report_pptx(project,document,lambda _:IMAGE))
        assert not any(n.startswith('ppt/charts/chart') for n in z.namelist())
        assert '10.0(1)' in text


def test_long_sample_names_do_not_add_unrequested_pages():
    project,document=fixture()
    name='긴샘플이름'*30
    document['samples'][0]['display_name']=name
    _,slides,text=deck(render_report_pptx(project,document,lambda _:IMAGE))
    assert len(slides)==5 and '긴 제목 전체' not in text
    assert any(name in description for s in slides for description in s.xpath('.//p:cNvPr/@descr',namespaces=NS))
    assert document['samples'][0]['display_name']==name


def test_legacy_options_and_prebuilt_pages_do_not_append_extra_slides(monkeypatch):
    project,document=fixture()
    document['figure']['include']=True
    document['quantity_table']['include']=True
    document['phase_classes']['criteria_text']='보존할 분류 기준'
    before=deepcopy((project,document))
    from app.xrd_workbench.render import posco_pptx
    monkeypatch.setattr(posco_pptx,'build_pages',lambda *a,**k: (_ for _ in ()).throw(AssertionError('legacy layout called')))
    old_pages=[dict(kind=k,items=[]) for k in ('cover','figure','table','appendix')]
    for pages in (None,old_pages):
        _,slides,text=deck(render_report_pptx(project,document,lambda _:IMAGE,pages=pages))
        assert len(slides)==5
        assert '보존할 분류 기준' not in text
    assert before==(project,document)


def test_phase_names_in_embedded_workbook_are_literal_strings_not_formulas():
    output=Deck();slide=output.slide(blank=True)
    assert output.pie(slide,[{'phase':'=1+1','value':100,'raw':'100.0','unit':'wt%'}],(20,30,100,100))
    book=load_workbook(io.BytesIO(output.charts[0][1]),data_only=False)
    assert book.active['A2'].value=='=1+1'
    assert book.active['A2'].data_type=='s'


def test_wppf_is_once_per_sample_and_its_entire_phase_table_stays_with_the_graph():
    from app.xrd_workbench.render.posco_pptx import build_posco_deck
    project = make_project([sample('A', n_phases=8), sample('B', n_phases=4),
                            sample('C', n_phases=2, quantities=0)])
    document = new_document(project)
    document['figure']['include'] = False
    before = deepcopy((project, document))
    filled = build_posco_deck(project, document, lambda _: IMAGE, image_renderer=lambda *a: IMAGE)
    titles = [''.join(named(s,'Title')[0].xpath('.//a:t/text()',namespaces=NS)) for s,_ in filled.pages]
    assert titles.count(QUANTITY_TITLE) == 2
    assert QUANTITY_TITLE+' (계속)' not in titles
    assert all(t.removesuffix(' (계속)') in ('XRD 분석 보고서','분석 이력 요약',SUMMARY_TITLE, QUANTITY_TITLE, PHASE_TITLE) for t in titles)
    wppf_pages = [s for s,_ in filled.pages if 'WPPF' in ''.join(named(s,'Title')[0].itertext())]
    assert [filled.page_info[s]['sample_id'] for s in wppf_pages] == ['A','B']
    # Every phase table row appears exactly once, even with unequal phase counts.
    cells = [text for s,_ in filled.pages for t in named(s,'Phase table')
             for text in t.xpath('.//a:tr/a:tc[1]//a:t/text()',namespaces=NS)]
    assert cells.count('[미분류] Phase0') == 1
    assert cells.count('Phase0 [미분류]') == 2  # WPPF wording only; layout-only phase pages retain their labels
    assert cells.count('Phase3 [미분류]') == 2
    assert cells.count('Phase7 [미분류]') == 1
    assert '이전 페이지에 수록' not in cells
    assert [len(named(s,'Phase table')[0].findall('.//a:tr',NS))-1 for s in wppf_pages] == [8,4]
    assert not any(named(s,'Phase continuation') for s,_ in filled.pages)
    assert before == (project, document)


def test_phase_table_uses_real_common_database_and_falls_back_for_mixed_or_missing():
    from app.xrd_workbench.render.posco_pptx import phase_table_data
    s = sample('db')
    for p in s['phases']: p['db'] = 'S/M:PDF-2 2025'
    header, rows = phase_table_data(s, lambda *a: None)
    assert header == 'S/M:PDF-2 2025'
    assert [r[2] for r in rows] == ['c0','c1','c2']
    s['phases'][1]['db'] = 'Other DB'
    header, rows = phase_table_data(s, lambda *a: None)
    assert header == 'DB / Card No.'
    assert rows[0][2] == 'S/M:PDF-2 2025 c0' and rows[1][2] == 'Other DB c1'
    s['phases'][1]['db'] = None
    assert phase_table_data(s, lambda *a: None)[0] == 'DB / Card No.'


def test_formula_subscripts_are_editable_and_preserve_raw_text_coefficients_and_charges():
    from app.xrd_workbench.render.posco_pptx import fit_text, subscript_formula
    slide = Deck().templates[1]
    prototype = named(slide,'Phase table')[0].find('.//a:tr[2]/a:tc[2]',NS)
    for formula, expected in [('Zn0.95 Mn0.05 O', ['0.95','0.05']),
                              ('Al2(SO4)3', ['2','4','3']), ('CuSO4·5H2O', ['4','2']),
                              ('Fe3+', []), ('13C', []), ('Li1-xNixO2', ['2'])]:
        cell = deepcopy(prototype)
        fit_text(cell,formula,dimensions=(80,12),size=9.5,bold=True)
        subscript_formula(cell)
        assert ''.join(cell.xpath('.//a:t/text()',namespaces=NS)) == formula
        assert cell.xpath('.//a:r[a:rPr/@baseline="-25000"]/a:t/text()',namespaces=NS) == expected
        assert all(v=='1' for v in cell.xpath('.//a:rPr/@b',namespaces=NS))
