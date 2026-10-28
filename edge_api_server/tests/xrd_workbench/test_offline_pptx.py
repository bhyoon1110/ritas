"""Read embedded files with an independent CFB reader, and follow actual PPT clicks."""
from copy import deepcopy
import io
import json
import posixpath
import re
import struct
from zipfile import ZipFile

from lxml import etree as E
import olefile
import pytest

from app.xrd_workbench.core.project import build_project
from app.xrd_workbench.core.report import new_document
from app.xrd_workbench.render import offline_pptx
from app.xrd_workbench.render.ole_package import package_file
from app.xrd_workbench.render.posco_pptx import render_report_pptx, NS
from app.xrd_workbench.render.link_audit import ppt_click_targets
from .synthetic import sample_files


def unpack_file(blob):
    with olefile.OleFileIO(blob, raise_defects=olefile.DEFECT_INCORRECT) as ole:
        native = ole.openstream('\1Ole10Native').read()
        assert not ole.parsing_issues
    assert struct.unpack_from('<I', native)[0] == len(native)-4
    assert struct.unpack_from('<H', native, 4)[0] == 2
    offset = native.index(b'\0', 6)+1
    offset = native.index(b'\0', offset)+1+4
    length = struct.unpack_from('<I', native, offset)[0]
    offset += 4+length
    length = struct.unpack_from('<I', native, offset)[0]
    data = native[offset+4:offset+4+length]
    offset += 4+length
    names = []
    for _ in range(3):
        length = struct.unpack_from('<I', native, offset)[0]*2
        names.append(native[offset+4:offset+4+length].decode('utf-16le'))
        offset += 4+length
    assert offset == len(native) and names[0] == names[1] == names[2]
    return names[0], data


def embedded_files(content):
    with ZipFile(io.BytesIO(content)) as archive:
        return dict(unpack_file(archive.read(n)) for n in archive.namelist()
                    if n.startswith('ppt/embeddings/xrdAttachment') and n.endswith('.bin'))


@pytest.fixture(scope='module')
def inputs():
    originals = {**sample_files('SampleA'), **sample_files('SampleB')}
    project = build_project(list(originals.items()))
    records = {r['id']: r for r in project['sources']}
    return project, new_document(project), lambda sid: originals[records[sid]['filename']]


@pytest.fixture(scope='module')
def delivered(inputs):
    return render_report_pptx(*inputs, delivery=True, link_base='https://inaccessible.invalid', storage_id='private')


@pytest.mark.parametrize('size', [1, 200, 3900, 4100, 100000, 10*1024*1024])
def test_ole_small_large_streams_and_extended_fat_preserve_unicode_filename_and_bytes(size):
    data = bytes(range(256))*(size//256) + bytes(range(size % 256))
    assert unpack_file(package_file('시료 A #1 (50%) ﹙원본﹚.txt', data)) == ('시료 A #1 (50%) ﹙원본﹚.txt', data)


@pytest.mark.parametrize('name', ['../file.txt', 'C:\\file.pdf', 'payload.exe', 'bad\0.txt'])
def test_embedding_never_accepts_commands_paths_or_executables(name):
    with pytest.raises(ValueError): package_file(name, b'bytes')


def test_every_linked_original_and_each_independent_viewer_are_embedded(inputs, delivered):
    project, _, source = inputs
    files = embedded_files(delivered)
    for sample in project['samples']:
        for record in sample['sources']:
            if record['role'] in {'measurement', 'phase_report', 'wppf_report'}:
                assert files[record['filename']] == source(record['id'])
        html = files[sample['name']+'.html'].decode()
        payload = json.loads(re.search(r'<script id="xrd-data" type="application/json">(.*?)</script>', html, re.S)[1])
        assert [s['id'] for s in payload['project']['samples']] == [sample['id']]
        assert payload['project']['samples'][0]['measurement'] == sample['measurement']
        assert 'location.replace' not in html
    assert '전체 비교 그래프.html' in files
    assert not ppt_click_targets(delivered)


def test_office_object_model_recognizes_direct_embedded_buttons_without_slide_jumps(delivered):
    from pptx import Presentation
    from pptx.enum.action import PP_ACTION
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    counts = {'jumps':0, 'files':0}
    for slide in Presentation(io.BytesIO(delivered)).slides:
        for shape in slide.shapes:
            if shape.click_action.action == PP_ACTION.NAMED_SLIDE:
                assert shape.click_action.target_slide is not None
                counts['jumps'] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
                assert shape.click_action.action == PP_ACTION.OLE_VERB
                assert shape.ole_format.prog_id == 'Package'
                assert unpack_file(shape.ole_format.blob)[1]
                counts['files'] += 1
    assert counts['jumps'] == 0 and counts['files'] > 5


def test_missing_or_empty_embedded_file_fails_audit(delivered):
    with ZipFile(io.BytesIO(delivered)) as original:
        for empty in [False, True]:
            missing = 'ppt/embeddings/xrdAttachment1.bin'
            output = io.BytesIO()
            with ZipFile(output, 'w') as changed:
                for name in original.namelist():
                    if name != missing: changed.writestr(name, original.read(name))
                    elif empty: changed.writestr(name, b'')
            with pytest.raises(ValueError, match='PPT 내부'):
                ppt_click_targets(output.getvalue())


def test_main_buttons_open_exact_files_in_place_without_adding_any_slides(inputs, delivered):
    native = render_report_pptx(*inputs)
    from app.xrd_workbench.render.link_audit import ppt_click_targets_by_slide
    expected = ppt_click_targets_by_slide(native)
    with ZipFile(io.BytesIO(delivered)) as z, ZipFile(io.BytesIO(native)) as source:
        source_parts = {n for n in source.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)}
        assert {n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)} == source_parts
        assert z.read('ppt/presentation.xml') == source.read('ppt/presentation.xml')
        assert not any('xrd-appendix' in n for n in z.namelist())
        assert len([n for n in z.namelist() if n.startswith('ppt/embeddings/xrdAttachment')]) == len({target for values in expected.values() for target in values})
        size = E.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz', NS)
        width, height = int(size.get('cx')), int(size.get('cy'))
        for path, old in expected.items():
            slide = E.fromstring(z.read(path))
            assert slide.get('show', '1') == '1'
            rels = {r.get('Id'):r for r in E.fromstring(z.read('ppt/slides/_rels/'+posixpath.basename(path)+'.rels'))}
            clicks = slide.findall('.//a:hlinkClick', NS)
            assert len(clicks) == len(old)
            originals = E.fromstring(source.read(path)).findall('.//a:hlinkClick', NS)
            for click, target in zip(clicks, old):
                assert click.get('action') == 'ppaction://ole?verb=0'
                assert click.get('{'+NS['r']+'}id') == ''
                frame = click.getparent().getparent().getparent()
                assert frame.tag == '{'+NS['p']+'}graphicFrame'
                obj = frame.find('.//p:oleObj', NS)
                rel = rels[obj.get('{'+NS['r']+'}id')]
                assert rel.get('Type') == NS['r']+'/oleObject' and rel.get('TargetMode') is None
                attachment = rel.get('Target')
                name, data = unpack_file(z.read(posixpath.normpath('ppt/slides/'+attachment)))
                from urllib.parse import unquote
                if target.startswith(('Report/', 'Measured%20Data/')):
                    assert name == unquote(target).rsplit('/',1)[-1]
                elif target.startswith('Viewer/'):
                    assert name == unquote(target).rsplit('/',1)[-1]
                else:
                    assert name == '전체 비교 그래프.html'
                assert data
            for click, original in zip(clicks, originals):
                frame = click.getparent().getparent().getparent()
                original_shape = original
                while original_shape.getparent().tag != '{'+NS['p']+'}spTree': original_shape = original_shape.getparent()
                old_xf = original_shape.find('p:spPr/a:xfrm', NS)
                if old_xf is None: old_xf = original_shape.find('p:xfrm', NS)
                xf = frame.find('p:xfrm', NS)
                start, extent = xf.find('a:off', NS), xf.find('a:ext', NS)
                assert dict(start.attrib) == dict(old_xf.find('a:off', NS).attrib)
                assert dict(extent.attrib) == dict(old_xf.find('a:ext', NS).attrib)
                assert 0 <= int(start.get('x')) <= width-int(extent.get('cx'))+1
                assert 0 <= int(start.get('y')) <= height-int(extent.get('cy'))+1


def test_excluded_samples_and_detached_output_have_no_hidden_data(inputs):
    project, original, source = inputs
    doc = deepcopy(original); doc['samples'][1]['included'] = False
    files = embedded_files(render_report_pptx(project, doc, source, delivery=True))
    assert all('SampleB' not in name for name in files)
    html = files['전체 비교 그래프.html'].decode()
    data = json.loads(re.search(r'<script id="xrd-data" type="application/json">(.*?)</script>', html, re.S)[1])
    assert len(data['project']['samples']) == 1
    assert embedded_files(render_report_pptx(project, doc, source, delivery=True, include_sources=False)) == {}


def test_oversize_attachment_fails_explicitly_instead_of_dropping_a_file(inputs, monkeypatch):
    monkeypatch.setattr(offline_pptx, 'MAX_BYTES', 1)
    with pytest.raises(ValueError, match='128MB'):
        render_report_pptx(*inputs, delivery=True)
