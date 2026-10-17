import io

from .figure_fixtures import make_document, make_project, make_sample
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from app.xrd_workbench.core.report import new_document
from app.xrd_workbench.render.html import render_html
from app.xrd_workbench.render.layout import BOTTOM, build_pages
from app.xrd_workbench.render.pdf import render_pdf
from app.xrd_workbench.render.pptx import render_pptx

A = ('PhaseA', '00-001', 'DB1')
B = ('PhaseB', '00-002', 'DB1')


def sample(i, phases=(A,)):
    return make_sample(f's{i}', f'시료 {i}', list(phases), [(1, 21.0, [p[0] for p in phases])])


def pages_for(samples, **figure):
    project = make_project(samples)
    document = make_document(project, **figure)
    return build_pages(project, document, lambda fid: b''), document


def figure_pages(pages):
    return [p for p in pages if p['kind'] == 'figure']


def texts(page):
    return ' '.join(i['text'] for i in page['items'] if i['type'] == 'text')


def test_figure_page_is_the_second_slide_after_the_cover():
    pages, _ = pages_for([sample(1), sample(2)])
    assert [p['kind'] for p in pages][:3] == ['cover', 'figure', 'summary']
    image = [i for i in figure_pages(pages)[0]['items'] if i['type'] == 'image']
    assert len(image) == 1 and image[0]['data'][:4] == b'\x89PNG'


def test_old_documents_without_a_figure_keep_their_pages():
    project = make_project([sample(1), sample(2)])
    document = new_document(project)
    del document['figure']
    del document['quantity_table']
    pages = build_pages(project, document, lambda fid: b'')
    assert [p['kind'] for p in pages] == ['cover', 'summary', 'sample', 'sample']


def test_include_off_adds_no_figure_page():
    pages, _ = pages_for([sample(1)], include=False)
    assert figure_pages(pages) == []


def test_every_sample_is_drawn_once_across_continued_pages_for_thirty_samples():
    pages, _ = pages_for([sample(i) for i in range(30)])
    fp = figure_pages(pages)
    assert len(fp) > 1
    assert fp[0]['title'] == '시료 비교 그래프' and fp[1]['title'].endswith('(계속)')
    assert all(i['y'] + i['h'] <= BOTTOM for p in fp for i in p['items'] if not i.get('role'))
    assert sum(len([i for i in p['items'] if i['type'] == 'image']) for p in fp) == len(fp)


def test_all_excluded_prints_the_reasons_instead_of_an_empty_figure():
    bad = make_sample('s1', '시료 1', [A], [], unit=None)
    pages, _ = pages_for([bad])
    fp = figure_pages(pages)
    assert len(fp) == 1
    assert not [i for i in fp[0]['items'] if i['type'] == 'image']
    assert '시료 1' in texts(fp[0]) and '축 또는 강도 단위 미확인' in texts(fp[0])


def test_excluded_samples_are_listed_under_the_figure():
    ok, bad = sample(1), make_sample('s2', '시료 2', [A], [], unit=None)
    pages, _ = pages_for([ok, bad])
    assert '제외: 시료 2' in texts(figure_pages(pages)[0])
    assert '축 또는 강도 단위 미확인' in texts(figure_pages(pages)[0])


def test_merge_rule_and_caption_are_printed():
    pages, _ = pages_for([sample(1)], caption='계열 비교')
    text = texts(figure_pages(pages)[0])
    assert '계열 비교' in text and '상 이름·Card 번호·데이터베이스가 모두 같은 상만' in text


def test_each_group_gets_its_own_figure_page():
    pages, _ = pages_for(
        [sample(1), sample(2, (B,)), sample(3)],
        groups=[{'title': '계열 1', 'samples': ['s1', 's3']}],
    )
    titles = [p['title'] for p in figure_pages(pages)]
    assert titles == ['시료 비교 그래프 · 계열 1', '시료 비교 그래프 · 기타']


def test_too_many_phase_rows_for_one_page_drop_the_reference_rows_with_a_note():
    many = tuple((f'Phase{i}', f'c{i}', 'DB') for i in range(60))
    project = make_project([make_sample('s1', '시료 1', list(many), [(1, 21.0, [many[0][0]])])])
    pages = build_pages(project, make_document(project), lambda fid: b'')
    fp = figure_pages(pages)
    assert [i for p in fp for i in p['items'] if i['type'] == 'image']
    assert '기준 피크 줄이 많아' in ' '.join(texts(p) for p in fp)
    assert all(i['y'] + i['h'] <= BOTTOM for p in fp for i in p['items'] if not i.get('role'))


def test_pdf_html_and_pptx_carry_the_figure():
    pages, document = pages_for([sample(1), sample(2)])
    assert len(PdfReader(io.BytesIO(render_pdf(pages, document))).pages) == len(pages)
    assert 'data:image/png' in render_html(pages, document, 'package')
    slides = Presentation(io.BytesIO(render_pptx(pages, document))).slides
    assert len(slides) == len(pages)
    kinds = [s.shape_type for s in slides[1].shapes]
    assert MSO_SHAPE_TYPE.PICTURE in kinds


def test_continued_pages_draw_each_sample_exactly_once_in_order(monkeypatch):
    import app.xrd_workbench.render.layout as layout

    seen = []
    real = layout.render_panel_png

    def spy(chunk, width, height, row_mm=5.0):
        seen.append([c['sample_id'] for c in chunk['curves']])
        return real(chunk, width, height, row_mm)

    monkeypatch.setattr(layout, 'render_panel_png', spy)
    pages_for([sample(i) for i in range(30)])
    assert [i for chunk in seen for i in chunk] == [f's{i}' for i in range(30)]
    assert len(seen) > 1


def test_merge_note_is_not_printed_when_no_phase_is_drawn():
    pages, _ = pages_for([sample(1)], show_markers=False, show_reference=False)
    assert '이름·Card 번호' not in texts(figure_pages(pages)[0])


def test_a_long_exclusion_list_continues_on_more_pages_without_clipping():
    bad = [make_sample(f'x{i}', f'제외 시료 {i}', [A], [], unit=None) for i in range(300)]
    pages, _ = pages_for([sample(1), *bad])
    fp = figure_pages(pages)
    assert len(fp) > 1
    assert all(i['y'] + i['h'] <= BOTTOM for p in fp for i in p['items'] if not i.get('role'))
    # Wrapping trims spaces at line ends, so compare with all whitespace removed.
    flat = ''.join(
        i['text'] for p in fp for i in p['items'] if i['type'] == 'text' and not i.get('role')
    )
    flat = ''.join(flat.split())
    assert all(f'제외시료{i}(축또는강도단위미확인)' in flat for i in range(300))


def test_nothing_drawable_with_many_excluded_samples_still_fits_the_pages():
    bad = [make_sample(f'x{i}', f'제외 시료 {i}', [A], [], unit=None) for i in range(300)]
    pages, _ = pages_for(bad)
    fp = figure_pages(pages)
    assert len(fp) > 1
    assert all(i['y'] + i['h'] <= BOTTOM for p in fp for i in p['items'] if not i.get('role'))


def test_a_few_curves_are_drawn_taller_up_to_the_cap_instead_of_leaving_the_page_empty():
    from app.xrd_workbench.render.layout import AXIS_MM, CURVE_MAX_MM, CURVE_MM, REF_MM

    pages, _ = pages_for([sample(1), sample(2)])
    image = next(i for i in figure_pages(pages)[0]['items'] if i['type'] == 'image')
    assert image['h'] == AXIS_MM + 2 * CURVE_MAX_MM + 1 * REF_MM
    assert CURVE_MAX_MM > CURVE_MM
    # Many curves keep the minimum height so the figure still fits and continues on more pages.
    pages, _ = pages_for([sample(i) for i in range(30)])
    first = next(i for i in figure_pages(pages)[0]['items'] if i['type'] == 'image')
    assert first['y'] + first['h'] <= BOTTOM


def spy_chunks(monkeypatch):
    import app.xrd_workbench.render.layout as layout

    calls = []
    real = layout.render_panel_png

    def spy(chunk, width, height, row_mm=5.0):
        calls.append((len(chunk['curves']), height, row_mm))
        return real(chunk, width, height, row_mm)

    monkeypatch.setattr(layout, 'render_panel_png', spy)
    return calls


def test_every_page_of_one_comparison_draws_curves_at_the_same_height(monkeypatch):
    from app.xrd_workbench.render.layout import AXIS_MM

    calls = spy_chunks(monkeypatch)
    pages_for([sample(i) for i in range(11)])
    assert len(calls) > 1
    heights = {round((h - AXIS_MM - 1 * row_mm) / n, 3) for n, h, row_mm in calls}
    assert len(heights) == 1, heights


def test_an_exclusion_line_stays_on_the_figure_page_when_curves_are_stretched():
    for curves in (2, 5, 6, 8):
        bad = make_sample('x1', '제외 시료', [A], [], unit=None)
        pages, _ = pages_for([*[sample(i) for i in range(curves)], bad])
        fp = figure_pages(pages)
        assert len(fp) == 1, (curves, len(fp))
        assert '제외: 제외 시료' in texts(fp[0])
