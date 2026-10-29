from copy import deepcopy
import json
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from tests.test_mail_address_book import client
from tests.test_ppt_templates import HEADERS, URL
from tests.test_tem_report_editor import report
from tests.xrd_workbench.test_package_names import collision_project
from app import ppt_templates
from app.report import ppt_template_import as importer, canvas, renderers
from app.report.model import ReportDocument
from app.report.direct_edit import edit_pages


@pytest.fixture(scope='module')
def source(tmp_path_factory):
    root=tmp_path_factory.mktemp('ppt-import'); path=root/'source.pptx'
    ppt=Presentation();ppt.slide_width=Inches(12);ppt.slide_height=Inches(6.75)
    ppt.slide_master.background.fill.solid();ppt.slide_master.background.fill.fore_color.rgb=RGBColor(221,238,255)
    slide=ppt.slides.add_slide(ppt.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(1),Inches(.8),Inches(9),Inches(.7))
    title.name='Template title';run=title.text_frame.paragraphs[0].add_run();run.text='OLD ANALYSIS NOT TO COPY';run.font.size=Pt(28);run.font.color.rgb=RGBColor(17,85,119)
    table=slide.shapes.add_table(2,2,Inches(1),Inches(2),Inches(6),Inches(2)).table
    table.cell(0,0).text='OLD VALUE';table.cell(1,0).text='OLD SAMPLE'
    ppt.save(path)
    return path


def upload(client, source, revision=0, name='외부 양식', experiment='FTIR'):
    return client.post(URL+'/import',params=dict(name=name,experiment=experiment,revision=revision),content=source.read_bytes(),headers=HEADERS)


@pytest.fixture(scope='module')
def prepared(source,tmp_path_factory):
    root=tmp_path_factory.mktemp('prepared')
    metadata=importer.prepare(source,root)
    return root,metadata


@pytest.fixture
def cached_prepare(monkeypatch,prepared):
    import shutil
    root,metadata=prepared
    def prepare(source,destination):
        for p in root.iterdir():shutil.copyfile(p,destination/p.name)
        return deepcopy(metadata)
    monkeypatch.setattr(importer,'prepare',prepare)


def test_real_conversion_extracts_bounded_styles_and_background_without_old_values(prepared):
    root,metadata=prepared
    assert metadata['ratio']==16/9
    objects=metadata['pages'][0]['objects']
    title=next(o for o in objects if o['kind']=='text')
    assert title['style']['fontSize']==28 and title['style']['color']=='#115577'
    assert 'OLD' in title['label'] # Private mapping preview only.
    assert 'text' not in json.dumps(title['style'])
    assert all('text' not in c for o in objects for c in o['cells'])
    with Image.open(root/'background-0.png') as image:
        assert image.getpixel((500,300))==(221,238,255)
        assert len(image.getcolors(image.width*image.height))==1
    assert not list(root.glob('*.pptx')) and not list(root.glob('*.pdf'))


def test_inherited_theme_properties_do_not_replace_current_report_styles():
    from lxml import etree as E
    shape=E.fromstring(f'<p:sp xmlns:p="{canvas.NS["p"]}" xmlns:a="{canvas.NS["a"]}"><p:spPr/><p:txBody><a:p><a:r><a:rPr sz="2800"/><a:t>Title</a:t></a:r></a:p></p:txBody></p:sp>')
    styles=importer.explicit_style(shape,dict(x=.1,y=.2,w=.3,h=.4,fontSize=28,color='#172B4D',bold=False,fill='none',stroke='none',align='left',lineWidth=1,dash='solid'))
    assert styles==dict(x=.1,y=.2,w=.3,h=.4,fontSize=28)


def test_import_private_mapping_revision_and_asset_lifecycle(client,source,cached_prepare,tmp_path):
    response=upload(client,source);assert response.status_code==201,response.text
    item=response.json()['selectedId'];url=URL+'/'+item
    result=client.get(url,headers=HEADERS).json();data=result['content']
    assert data['version']==2 and data['bindings']==[]
    data['bindings']=[dict(key='a'*64,slide=0,background=True,objects=[dict(source='slot-0',target='shape-'+'b'*20)])]
    response=client.put(url,json=dict(revision=1,content=data),headers=HEADERS);assert response.status_code==200,response.text
    assert client.put(url,json=dict(revision=1,content=data),headers=HEADERS).status_code==409
    assert client.get(url+'/source',headers=HEADERS).content==source.read_bytes()
    background=client.get(url+'/slides/0/background',headers=HEADERS)
    assert background.status_code==200 and 'no-store' in background.headers['cache-control']
    assert client.get(url+'/slides/12/preview',headers=HEADERS).status_code==404
    for tail in ('','/source','/slides/0/background'):
        assert client.get(url+tail,headers={**HEADERS,'x-test-user':'two'}).status_code==404
    job_root=tmp_path/'report-assets';asset=canvas.store_image(job_root,background.content)['asset']
    source_root=next((tmp_path/'user-ppt-templates').glob('*.json')).with_suffix('')/data['source']
    response=client.request('DELETE',url,json=dict(revision=2),headers=HEADERS)
    assert response.status_code==200 and not source_root.exists()
    assert (job_root/(asset+'.png')).is_file()


def test_import_limits_origin_missing_source_and_stale_input(client,source,cached_prepare,monkeypatch,tmp_path):
    assert client.post(URL+'/import',params=dict(name='a',experiment='FTIR',revision=0),content=b'x',headers={}).status_code==401
    assert client.post(URL+'/import',params=dict(name='a',experiment='FTIR',revision=0),content=b'x',headers={**HEADERS,'Origin':'https://other.invalid'}).status_code==403
    monkeypatch.setattr(importer,'MAX_BYTES',10)
    assert upload(client,source).status_code==413
    monkeypatch.setattr(importer,'MAX_BYTES',10*1024*1024)
    assert upload(client,source,1).status_code==409
    assert upload(client,source).status_code==201
    response=upload(client,source,1,name='duplicate source');assert response.status_code==201
    item=response.json()['selectedId'];data=client.get(URL+'/'+item,headers=HEADERS).json()['content']
    bad=deepcopy(data);bad['source']='0'*64
    assert client.put(URL+'/'+item,json=dict(revision=2,content=bad),headers=HEADERS).status_code==410
    bad=deepcopy(data);bad['bindings']=[dict(key='a'*64,slide=0,background=True,objects=[dict(source='slot-99',target='shape-'+'b'*20)])]
    assert client.put(URL+'/'+item,json=dict(revision=2,content=bad),headers=HEADERS).status_code==422
    with ppt_templates._IMPORT_SLOT:
        assert upload(client,source,2,name='busy').status_code==409
    assert not list(tmp_path.rglob('.import-*'))


@pytest.mark.parametrize('kind',['macro','ole','svg','external','traversal','doctype','badzip'])
def test_unsafe_or_invalid_package_is_rejected_before_rendering(source,tmp_path,kind):
    with ZipFile(source) as z:files={i:z.read(i) for i in z.namelist()}
    if kind=='macro':files['ppt/vbaProject.bin']=b'no'
    elif kind=='ole':files['ppt/embeddings/ole.bin']=b'no'
    elif kind=='svg':files['ppt/media/image.svg']=b'<svg/>'
    elif kind=='doctype':files['ppt/slides/slide1.xml']=b'<!DOCTYPE a [<!ENTITY b SYSTEM "file:///etc/passwd">]><a>&b;</a>'
    elif kind in ('external','traversal'):
        name='ppt/slideMasters/_rels/slideMaster1.xml.rels'
        files[name]=b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId99" Type="image" Target="'+(b'https://invalid/image.png" TargetMode="External' if kind=='external' else b'../../../outside.png')+b'"/></Relationships>'
    path=tmp_path/'bad.pptx'
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for n,b in files.items():z.writestr(n,b)
    if kind=='badzip':path.write_bytes(b'not zip')
    with pytest.raises(ValueError):importer.read_package(path)


@pytest.mark.parametrize('experiment',['FT-IR','RAMAN'])
def test_applied_external_style_and_background_preserve_report_values(prepared,tmp_path,experiment):
    root,meta=prepared
    doc=ReportDocument('new','NEW ANALYSIS '+experiment,experiment,{},'2026-10-29')
    path=tmp_path/'output.pptx';renderers.render_pptx(doc,path)
    page=edit_pages(doc,path)['pages'][0];title=next(o for o in page['objects'] if o['name']=='Report title')
    style=next(o['style'] for o in meta['pages'][0]['objects'] if o['kind']=='text')
    asset=canvas.store_image(tmp_path/'assets',(root/'background-0.png').read_bytes())['asset'];doc.canvas_asset_root=str(tmp_path/'assets')
    doc.canvas_edits={page['id']:[dict(style,id=title['id'],kind='text'),dict(id='object-template-background',kind='image',asset=asset,x=0,y=0,w=1,h=1,layer='back')]}
    renderers.render_pptx(doc,path)
    final=edit_pages(doc,path)['pages'][0];changed=next(o for o in final['objects'] if o['id']==title['id'])
    assert changed['text']=='NEW ANALYSIS '+experiment and changed['fontSize']==28
    assert any(o['id']=='object-template-background' for o in final['objects'])
    with ZipFile(path) as z:
        assert b'OLD ANALYSIS' not in b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
        assert b'Template RIST logo' in z.read('ppt/slides/slide1.xml')
        assert b'Template POSCO background' in z.read('ppt/slides/slide1.xml')
    assert page['templateKey']==final['templateKey']


def test_tem_external_style_and_background_preserve_native_evidence(report,prepared,tmp_path):
    from types import SimpleNamespace
    from app.report import tem_canvas
    h=report;root,metadata=prepared
    model=h.client.get(h.url+'/editor-data').json()['document']
    pages=h.client.get(h.url+'/edit-layout').json()['pages']
    page=next(p for p in pages if any(o['kind']=='text' and o['text'] for o in p['objects']))
    target=next(o for o in page['objects'] if o['kind']=='text' and o['text'])
    job=SimpleNamespace(work_dir=tmp_path/'new-tem',pptx_path=tmp_path/'tem-applied.pptx',job_id='template-qa',parent_job_id=h.job.job_id)
    tem_canvas.preserve_base(job,h.job)
    asset=canvas.store_image(tem_canvas.asset_root(job),(root/'background-0.png').read_bytes())['asset']
    style=metadata['pages'][0]['objects'][0]['style']
    model['canvasEdits']={page['id']:[dict(style,id=target['id'],kind='text'),dict(id='object-template-background',kind='image',asset=asset,x=0,y=0,w=1,h=1,layer='back')]}
    tem_canvas.render(job,model)
    actual=next(p for p in tem_canvas.layout(job,model)['pages'] if p['id']==page['id'])
    changed=next(o for o in actual['objects'] if o['id']==target['id'])
    assert changed['text']==target['text'] and changed['fontSize']==28
    assert any(o['id']=='object-template-background' for o in actual['objects'])


def test_xrd_external_style_and_background_preserve_phase_tables(collision_project,prepared,tmp_path):
    from app.xrd_workbench.render.posco_pptx import build_posco_deck, render_report_pptx
    from app.xrd_workbench.render.posco_preview import deck_pages
    from io import BytesIO
    project,document,data=deepcopy(collision_project);root,metadata=prepared
    pages=deck_pages(build_posco_deck(project,document,data.__getitem__))
    page=next(p for p in pages if any(o['kind']=='table' for o in p['objects']))
    table=next(o for o in page['objects'] if o['kind']=='table')
    asset=canvas.store_image(tmp_path,(root/'background-0.png').read_bytes())['asset']
    style=next(o['style'] for o in metadata['pages'][0]['objects'] if o['kind']=='table')
    document['canvasEdits']={page['id']:[dict(style,id=table['id'],kind='table'),dict(id='object-template-background',kind='image',asset=asset,x=0,y=0,w=1,h=1,layer='back')]}
    final=deck_pages(build_posco_deck(project,document,data.__getitem__,asset_root=tmp_path))
    actual=next(p for p in final if p['id']==page['id'])
    changed=next(o for o in actual['objects'] if o['id']==table['id'])
    assert [c['text'] for c in changed['cells']]==[c['text'] for c in table['cells']]
    with ZipFile(BytesIO(render_report_pptx(project,document,data.__getitem__,asset_root=tmp_path))) as z:
        assert any(n.startswith('ppt/media/canvas-') for n in z.namelist())
        assert b'OLD VALUE' not in b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
