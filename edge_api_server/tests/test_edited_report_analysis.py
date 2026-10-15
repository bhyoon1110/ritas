from copy import deepcopy
from dataclasses import replace
import json
import zipfile

import httpx
import pytest

from app.config import Settings
from app.llm_client import LocalLlmClient
from app import preview_report
from app.report import annotator
from app.report.builders import get_builder
from app.report.edited_analysis import prepare_edited_report, edited_graph_spec, analysis_preview
from app.report.spectrum_viewer import viewer_html
from app.report.renderers import render_pptx, _report_readable_text
from edge_api_server.tests.test_ftir_web import TINY_PNG_DATA_URL, use_fake_pptx_pdf_converter


def edited_payload():
    return {
        "samples": [{"label": "old sample", "fileName": "sample.dpt", "pointCount": 3, "peakCount": 99,
                     "sample_group": "sample:2", "metadata": {"Resolution": "4 cm-1"}}],
        "findings": {"functional_groups": [{"group": "DELETED CHEMISTRY"}]},
        "rule_matches": [{"compound": "DELETED CHEMISTRY"}],
        "figure": {"data": [{"name": "STALE FIGURE"}], "layout": {}},
        "viewerFigure": {"data": [
            {"name": "Edited sample", "x": [1500, 1700, 1900], "y": [11, 13, 11], "mode": "lines",
             "meta": {"rist_sample_parent": True, "rist_sample_group": "sample:2", "rist_view_offset": 10}},
            {"name": "수정 C=O 후보", "x": [1700], "y": [13], "mode": "markers",
             "meta": {"rist_sample_group": "sample:2", "rist_view_offset": 10,
                      "rist_peak": {"x": 1690, "base_y": 999, "sample_group": "sample:2", "user": True,
                                    "label": "DELETED ASSIGNMENT", "assignments": [{"name": "DELETED ASSIGNMENT"}]}}},
            {"name": "DELETED PEAK", "visible": False, "x": [1200], "y": [13],
             "meta": {"rist_peak": {"x": 1200, "sample_group": "sample:2", "label": "DELETED PEAK"}}},
        ], "layout": {
            "xaxis": {"range": [1900, 1500]}, "yaxis": {"title": {"text": "Absorbance"}},
            "annotations": [{"text": "수정 C=O 후보", "x": 1700, "y": 14},
                            {"text": "산화 가능성 비교", "x": 1720, "y": 14, "name": "rist_text_box_text:1"},
                            {"text": "삭제 메모", "visible": False},
                            {"text": "I(1700)/I(1500) = 3.0", "name": "rist_raman_ratio:1"}],
            "shapes": [{"type": "line", "x0": 1700, "x1": 1700, "y0": 13, "y1": 14},
                       {"name": "rist_rect:1", "type": "rect", "x0": 1750, "x1": 1650, "y0": 12, "y1": 14},
                       {"type": "rect", "xref": "paper", "yref": "paper", "x0": 0, "x1": 1, "y0": 0, "y1": 1},
                       {"name": "rist_text_box:1", "type": "rect", "x0": 1720, "x1": 1800, "y0": 14, "y1": 15}],
            "meta": {"ristPeakLabels": [{"traceIndex": 1, "annotationIndex": 0, "shapeIndex": 0}]},
        }},
    }


def build_spec(payload, experiment="FT-IR"):
    prepared = prepare_edited_report(payload)
    builder = get_builder(experiment)
    job = {"job_id": "test", "request_number": "REQ-EDIT", "experiment_code": experiment,
           "equipment_code": "", "operator_id": "", "_generated_at": "2026-10-14"}
    items = [{"data": prepared}]
    return prepared, builder.build(job, items), edited_graph_spec(builder.llm_slots(job, items), prepared, experiment)


@pytest.mark.parametrize("experiment", ["FT-IR", "RAMAN"])
def test_final_edit_evidence_excludes_old_results_and_keeps_box_coordinates(experiment):
    payload = edited_payload()
    before = deepcopy(payload)
    prepared, document, spec = build_spec(payload, experiment)
    assert payload == before
    facts = spec.facts
    assert "DELETED" not in json.dumps(facts)
    assert "STALE FIGURE" not in json.dumps(facts)
    assert facts["visible_peak_count"] == 1
    assert facts["current_peaks"][0]["position_cm-1"] == 1700
    assert facts["current_peaks"][0]["intensity"] == 3
    assert facts["current_peaks"][0]["display_intensity"] == 13
    assert facts["current_peaks"][0]["manual"] is True
    assert prepared["samples"][0]["peakCount"] == 1
    assert prepared["samples"][0]["fileName"] == "sample.dpt"
    assert prepared["samples"][0]["metadata"] == {"Resolution": "4 cm-1"}
    assert facts["graph_notes"]["text_notes"][0]["text"] == "산화 가능성 비교"
    regions = facts["graph_notes"]["regions"]
    assert len(regions) == 2
    assert regions[0]["included_peak_ids"] == ["p1"] and regions[0]["peak_count"] == 1
    assert "peak_count" not in regions[1]  # Paper coordinates are not wavenumbers.
    assert "삭제 메모" not in json.dumps(facts, ensure_ascii=False)
    assert "DELETED CHEMISTRY" not in document.to_markdown()
    assert "전처리 단계 검출 피크 후보" not in document.to_markdown()
    assert "전처리 검출 피크 후보" not in document.to_markdown()


@pytest.mark.parametrize("experiment", ["FT-IR", "RAMAN"])
def test_zero_peaks_do_not_revive_original_analysis(experiment):
    payload = edited_payload()
    payload["viewerFigure"]["data"] = payload["viewerFigure"]["data"][:1]
    _, document, spec = build_spec(payload, experiment)
    assert spec.facts["visible_peak_count"] == 0
    assert spec.facts["current_peaks"] == []
    assert spec.facts["samples"][0]["peak_count"] == 0
    assert "0이면 피크 기반 해석을 할 수 없다고" in spec.system_prompt
    assert "DELETED" not in json.dumps(spec.facts)
    if experiment == "RAMAN":
        sample_section = next(section for section in document.sections if section.section_id == "raman_samples")
        assert sample_section.table.rows[0][-1] == "0"


def test_empty_legacy_snapshot_does_not_retain_old_peak_count():
    payload = edited_payload()
    payload["viewerFigure"]["data"] = []
    prepared, _, spec = build_spec(payload)
    assert prepared["samples"][0]["peakCount"] == 0
    assert spec.facts["samples"][0]["peak_count"] == 0


def test_many_peaks_are_fairly_bounded_and_omissions_are_explicit():
    payload = edited_payload()
    traces = payload["viewerFigure"]["data"]
    traces[:] = traces[:2]
    parent, peak = deepcopy(traces[0]), deepcopy(traces[1])
    parent["meta"]["rist_sample_group"] = "sample:3"
    parent["name"] = "Second sample"
    traces.append(parent)
    for group in ("sample:2", "sample:3"):
        for index in range(60):
            item = deepcopy(peak)
            item["x"] = [1000+index]
            item["meta"]["rist_peak"]["sample_group"] = group
            item["meta"]["rist_sample_group"] = group
            traces.append(item)
    _, _, spec = build_spec(payload)
    assert spec.facts["visible_peak_count"] == 121
    assert spec.facts["included_peak_count"] == 48
    assert spec.facts["omitted_peak_count"] == 73
    assert {p["sample"] for p in spec.facts["current_peaks"]} == {"Edited sample", "Second sample"}


@pytest.mark.parametrize("failure", [False, True])
@pytest.mark.parametrize("experiment", ["FT-IR", "RAMAN"])
def test_llm_request_to_ppt_and_html_preview(tmp_path, monkeypatch, failure, experiment):
    use_fake_pptx_pdf_converter(monkeypatch)
    captured = []
    def handler(request):
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "gemma4-e4b", "max_model_len": 8192}]})
        body = json.loads(request.content)
        captured.append(body)
        if failure:
            return httpx.Response(503, json={"error": {"message": "MOCK_PRIVATE_LLM_FAILURE_503"}})
        response = {key: "1700 cm⁻¹의 수정 C=O 후보와 강조 영역을 확인했습니다. 사용자 메모의 산화 가능성은 추가 검토가 필요합니다."
                    for key in ("summary", "key_findings", "interpretation", "qc_notes", "narrative", "caption", "email_body", "email_subject")}
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(response, ensure_ascii=False)}}]})
    monkeypatch.setattr(preview_report, "LocalLlmClient", lambda *args: LocalLlmClient(*args, transport=httpx.MockTransport(handler)))
    settings = Settings(storage_root=tmp_path, llm_include_images=False)
    root, package = preview_report.build_preview_report_package(
        experiment_code=experiment, analysis_payload=edited_payload(),
        raw_series=[preview_report.RawSeries("sample", [1500, 1700, 1900], [1, 3, 1])],
        figure_image=preview_report.decode_figure_image(TINY_PNG_DATA_URL), settings=settings)
    try:
        assert captured
        prompt = captured[0]["messages"][-1]["content"]
        assert "DELETED" not in prompt and "STALE FIGURE" not in prompt
        assert '"position_cm-1":1700.0' in prompt and '"included_peak_ids":["p1"]' in prompt
        with zipfile.ZipFile(package) as archive:
            html = archive.read("report.html").decode()
            assert "분석 결과" in html and "로컬 AI" in html
            if failure:
                assert "로컬 AI 해석 미생성 · 규칙 기반 요약" in html
                assert "MOCK_PRIVATE_LLM_FAILURE_503" not in html
            else:
                assert "로컬 AI 해석 · 분석자 검토 필요" in html
                assert "사용자 메모의 산화 가능성은 추가 검토" in html
                from io import BytesIO
                with zipfile.ZipFile(BytesIO(archive.read("report.pptx"))) as ppt:
                    text = "".join(ppt.read(name).decode() for name in ppt.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
                assert "사용자 메모의 산화 가능성은 추가 검토" in text
    finally:
        preview_report.cleanup_preview_report(root)


def test_partial_output_is_not_reported_as_full_ai_and_html_is_escaped():
    _, document, _ = build_spec(edited_payload())
    document.apply_llm_slots({"summary": "<script>alert(1)</script>"})
    analysis = analysis_preview(document)
    assert analysis["status"] == "partial"
    html = viewer_html(title="Test", figure={}, slides=["data:image/jpeg;base64,a"], analysis=analysis)
    assert "<script>alert(1)" not in html and "&lt;script&gt;" in html
    assert "로컬 AI 해석 일부 생성" in html


def test_small_model_budget_reduces_facts_without_mutating_original(tmp_path):
    _, _, spec = build_spec(edited_payload())
    facts = deepcopy(spec.facts)
    facts["current_peaks"] = [{**facts["current_peaks"][0], "id": "p"+str(i), "label": "긴 피크 이름 "*30} for i in range(48)]
    facts.update(visible_peak_count=48, included_peak_count=48, omitted_peak_count=0)
    spec = replace(spec, facts=facts)
    class Client:
        def get_model_info(self):
            return {"max_model_len": 8192}
    content, tokens = annotator._fit_edited_evidence(Settings(storage_root=tmp_path, llm_include_images=False), Client(), spec, tmp_path)
    reduced = json.loads(content.split("\n\n", 1)[1])
    assert reduced["included_peak_count"] < 48 and reduced["omitted_peak_count"] > 0
    assert len(spec.facts["current_peaks"]) == 48 and tokens >= 768


@pytest.mark.parametrize("experiment", ["FT-IR", "RAMAN"])
def test_late_user_edits_take_priority_over_earlier_detected_peaks(experiment, tmp_path):
    payload = edited_payload()
    figure = payload["viewerFigure"]
    figure["data"] = figure["data"][:1]
    figure["layout"]["annotations"] = []
    figure["layout"]["shapes"] = [{"type": "rect", "x0": 1603, "x1": 1603, "y0": 12, "y1": 14}]
    figure["layout"]["meta"] = {}
    for i in range(104):
        peak = {"x": 1500 + i, "sample_group": "sample:2", "label": f"피크 {i}"}
        if i == 100:
            peak["user"] = True
        if i == 102:
            peak["group_name"] = "관심 그룹"
        figure["data"].append({"name": "최종 수정명" if i == 101 else peak["label"], "x": [1500+i], "y": [13],
                               "meta": {"rist_peak": peak, "rist_view_offset": 10}})
    _, _, spec = build_spec(payload, experiment)
    current = spec.facts["current_peaks"]
    assert [p["position_cm-1"] for p in current[:4]] == [1600, 1601, 1602, 1603]
    assert [p["edit_roles"] for p in current[:4]] == [["manual"], ["renamed"], ["grouped"], ["region"]]
    assert current[1]["label"] == "최종 수정명"
    assert spec.facts["graph_notes"]["regions"][0]["included_peak_ids"] == ["p104"]
    assert spec.facts["omitted_peak_count"] == 56
    class Client:
        def get_model_info(self):
            return {"max_model_len": 8192}
    content, _ = annotator._fit_edited_evidence(Settings(storage_root=tmp_path, llm_include_images=False,
                                                       llm_max_input_chars=4000), Client(), spec, tmp_path)
    reduced = json.loads(content.split("\n\n", 1)[1])
    assert reduced["included_peak_count"] < 48
    assert reduced["current_peaks"][0]["position_cm-1"] == 1600
    assert all(p["intensity"] == 3 for p in reduced["current_peaks"])


def test_ratio_evidence_uses_same_current_values_as_report_table():
    from edge_api_server.tests.test_raman_report_tables import payload_for_tables
    payload = payload_for_tables()
    payload["viewerFigure"]["data"][1]["name"] = "편집 D"
    payload["viewerFigure"]["data"][1]["y"] = [13]
    # Annotation is deliberately stale; final measurements are authoritative.
    payload["viewerFigure"]["layout"]["annotations"][0]["text"] = "STALE_RATIO=999"
    _, document, spec = build_spec(payload, "RAMAN")
    ratios = spec.facts["graph_notes"]["ratio_measurements"]
    assert len(ratios) == 2
    assert ratios[0]["sample"] == "시료 1" and ratios[1]["sample"] == "시료 2"
    assert ratios[0]["numerator"]["label"] == "편집 D"
    assert ratios[0]["numerator"]["intensity"] == 3
    assert ratios[0]["denominator"]["intensity"] == 4
    assert ratios[0]["value"] == 0.75
    assert ratios[0]["value_text"] == document.section("raman_ratios").table.rows[0][-1] == "0.750"
    assert "STALE_RATIO" not in json.dumps(spec.facts)
    assert "STALE_RATIO" not in document.to_markdown()  # Offline/failure fallback agrees too.
    assert all("ratio" in peak["edit_roles"] for peak in spec.facts["current_peaks"])


@pytest.mark.parametrize("kind", ["delete", "hide", "zero"])
def test_deleted_or_invalid_ratio_evidence_is_not_recreated_from_old_annotation(kind):
    from edge_api_server.tests.test_raman_report_tables import payload_for_tables
    payload = payload_for_tables(1)
    traces = payload["viewerFigure"]["data"]
    if kind == "delete":
        del traces[2]
    elif kind == "hide":
        traces[2]["visible"] = False
    else:
        traces[2]["y"] = [10]  # Zero after removing the display offset.
    _, _, spec = build_spec(payload, "RAMAN")
    notes = spec.facts["graph_notes"]
    assert notes["ratio_annotations"] == []
    if kind == "zero":
        assert notes["ratio_measurements"][0]["value"] is None
        assert notes["ratio_measurements"][0]["value_text"] == "계산 불가"
    else:
        assert notes["ratio_measurements"] == []


def test_small_context_keeps_zero_peak_notes_instead_of_failing(tmp_path):
    payload = edited_payload()
    payload["viewerFigure"]["data"] = payload["viewerFigure"]["data"][:1]
    payload["viewerFigure"]["layout"]["meta"] = {}
    payload["viewerFigure"]["layout"]["annotations"] = [{"text": f"메모 {i}: " + "검토 근거 " * 30} for i in range(8)]
    _, _, spec = build_spec(payload)
    class Client:
        def get_model_info(self):
            return {"max_model_len": 8192}
    content, _ = annotator._fit_edited_evidence(Settings(storage_root=tmp_path, llm_include_images=False,
                                                       llm_max_input_chars=1800), Client(), spec, tmp_path)
    reduced = json.loads(content.split("\n\n", 1)[1])
    assert reduced["current_peaks"] == [] and reduced["visible_peak_count"] == 0
    notes = reduced["graph_notes"]
    assert 0 < len(notes["text_notes"]) < 8
    assert len(notes["text_notes"]) + notes["omitted_notes"] == 8
    assert len(spec.facts["graph_notes"]["text_notes"]) == 8


def test_ppt_keeps_all_ai_caveats_when_summary_cards_shorten_them(tmp_path):
    from pptx import Presentation
    _, document, _ = build_spec(edited_payload())
    paragraph = "첫 번째 관찰입니다. 두 번째 검토사항입니다. 세 번째는 생략되면 안 되는 해석 한계입니다."
    long_text = "관심 영역의 피크는 후보 소견이며 반복 측정과 보완 분석이 필요합니다. " * 30 + "마지막 검토사항입니다."
    document.apply_llm_slots({"qc_notes": paragraph, "interpretation": long_text})
    target = tmp_path / "report.pptx"
    render_pptx(document, target)
    ppt = Presentation(target)
    text = "".join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
    normalized = "".join(text.split())
    for original in (paragraph, long_text):
        expected = "".join(_report_readable_text(original).split())
        # Long text can cross pages with headers; every sentence must still appear.
        assert all(sentence in normalized for sentence in expected.split(".") if sentence)
    assert "품질 확인 및 검토사항 · 전문" in text and "피크 해석 · 전문 · 2" in text
