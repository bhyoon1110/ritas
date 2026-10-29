from datetime import datetime, timezone
from io import BytesIO
import base64
import hashlib
from types import SimpleNamespace
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
import pytest

from app.auth import AuthContext
from app.report.pptx_protection import P, create_protection, protect_pptx, validate_protection
from tests.test_report_editor import prepared, edit_request, wait_revision
from tests.test_tem_report_editor import report, editable, wait

HEADERS = {'X-Requested-With': 'RIST-Report-Editor'}


def verifier(ppt):
    with ZipFile(BytesIO(ppt)) as archive:
        node = etree.fromstring(archive.read('ppt/presentation.xml')).find(f'{{{P}}}modifyVerifier')
        return dict(node.attrib) if node is not None else {}


def matches(password, value):
    # Independent reader of the serialized OOXML attributes, following ECMA/POI.
    result = hashlib.new('sha512', base64.b64decode(value['saltData']) + password.encode('utf-16le')).digest()
    for counter in range(int(value['spinCount'])):
        result = hashlib.new('sha512', result + counter.to_bytes(4, 'little')).digest()
    return result == base64.b64decode(value['hashData'])


@pytest.mark.parametrize('password', ['Test-123!', ' 한글 🔒 변경 암호 ', 'a'*255])
def test_office_modify_verifier_preserves_parts_links_and_native_objects(password):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(100, 100, 3000000, 1000000)
    run = box.text_frame.paragraphs[0].add_run(); run.text = 'Original text'
    run.hyperlink.address = 'Viewer/sample.html'
    data = BytesIO(); deck.save(data)
    original = data.getvalue()
    settings = create_protection(password)
    output = protect_pptx(original, settings)
    assert verifier(output) == settings
    assert matches(password, settings) and not matches('wrong password', settings)
    assert create_protection(password) != settings  # Fresh random salt per setting.
    assert not verifier(original)
    assert Presentation(BytesIO(output)).slides[0].shapes[0].text == 'Original text'
    with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(output)) as b:
        assert set(a.namelist()) == set(b.namelist())
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name != 'ppt/presentation.xml')
    replaced = protect_pptx(output, create_protection('changed'))
    assert matches('changed', verifier(replaced)) and not matches(password, verifier(replaced))


@pytest.mark.parametrize('password', ['', 'a'*256, '\ufeffsecret', '\ud800', None, 123])
def test_invalid_password_rejected_without_echo(password):
    with pytest.raises(ValueError) as error:
        create_protection(password)
    assert 'secret' not in str(error.value)


def test_verifier_rejects_arbitrary_attributes_spin_counts_and_invalid_base64():
    good = create_protection('test')
    for bad in [dict(good, password='secret'), dict(good, spinCount='999999999'),
                dict(good, hashData='x'*88), dict(good, saltData='AAAA'), [], 'secret']:
        with pytest.raises(ValueError): validate_protection(bad)


def test_spectrum_setting_survives_revision_download_reopen_and_removal(prepared):
    h = prepared
    original = h.job.package_path.read_bytes()
    endpoint = h.url + '/pptx-password'
    assert h.client.post(endpoint, json={'password':'secret'}).status_code == 403
    assert h.client.post(endpoint, json={'password':'secret'}, headers=dict(HEADERS, Origin='https://evil.example')).status_code == 403
    response = h.client.post(endpoint, json={'password':'분석 Secret 42'}, headers=HEADERS)
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    settings = response.json()['protection']
    assert 'Secret' not in response.text
    current = h.url
    for protection in [settings, None, {}]:
        value = edit_request(h.client.get(current+'/editor-data').json()['document'])
        if protection is None:
            value.pop('pptxProtection', None)  # Old clients preserve the setting.
        else:
            value['pptxProtection'] = protection
        result = h.client.post(current+'/revisions', json=value, headers=HEADERS)
        assert result.status_code == 202, result.text
        job_id = result.json()['jobId']
        assert wait_revision(h, job_id)['status'] == 'completed'
        job = h.store.get(job_id)
        current = h.prefix+'/'+job_id
        reopened = h.client.get(current+'/editor-data').json()['document']
        expected = {} if protection == {} else settings
        assert reopened['pptxProtection'] == expected
        with ZipFile(BytesIO(h.client.get(current+'/download').content)) as package:
            ppt = package.read('report.pptx')
            assert verifier(ppt) == expected
            assert any(shape.has_text_frame and shape.text
                       for slide in Presentation(BytesIO(ppt)).slides for shape in slide.shapes)
        # Preview/editing uses an unprotected authoring copy and no raw password.
        path = job.tmp_root/'job/report'
        assert not verifier((path/'report.pptx').read_bytes())
        assert 'Secret 42' not in (path/'editor-state.json').read_text()
        assert 'Secret 42' not in (path/'report.json').read_text()
    assert h.job.package_path.read_bytes() == original


def test_password_hash_endpoint_uses_report_ownership(prepared):
    h = prepared
    h.app.state.settings = SimpleNamespace(auth_enabled=True)
    h.app.state.actor = AuthContext(user_id='other',login_id='user',email=None,display_name='User',status='ACTIVE',
        session_id='test',session_expires_at=datetime.now(timezone.utc),sso_authenticated_at=None,
        projects=frozenset({h.kind.upper()}),roles=frozenset(),sso_identity=None)
    assert h.client.post(h.url+'/pptx-password',json={'password':'secret'},headers=HEADERS).status_code == 403


def test_tem_direct_ppt_and_lims_package_have_same_protection(report):
    from app import ahn_web
    h = report
    response = h.client.post(h.url+'/pptx-password', json={'password':'TEM-secret'}, headers=HEADERS)
    settings = response.json()['protection']
    current = h.url
    for protection in [settings, None, {}]:
        value = editable(h.client.get(current+'/editor-data').json()['document'])
        if protection is None: value.pop('pptxProtection', None)
        else: value['pptxProtection'] = protection
        result = h.client.post(current+'/revisions', json=value, headers=HEADERS)
        assert result.status_code == 202, result.text
        finished = wait(h, result.json())
        assert finished['status'] == 'completed', finished
        current = '/api/v1/tem/report/jobs/'+finished['jobId']
        expected = {} if protection == {} else settings
        direct = h.client.get(current+'/download/pptx').content
        with ZipFile(BytesIO(h.client.get(current+'/download').content)) as package:
            assert verifier(package.read('tem-report.pptx')) == expected
            # Separate serializations must still preserve all native content.
            with ZipFile(BytesIO(direct)) as a, ZipFile(BytesIO(package.read('tem-report.pptx'))) as b:
                assert {n:a.read(n) for n in a.namelist()} == {n:b.read(n) for n in b.namelist()}
        assert verifier(direct) == expected
        assert any(shape.has_text_frame and shape.text
                   for slide in Presentation(BytesIO(direct)).slides for shape in slide.shapes)
        job = ahn_web._ahn_report_jobs[finished['jobId']]
        assert not verifier(job.pptx_path.read_bytes())
        assert 'TEM-secret' not in (job.work_dir/'review/editor-state.json').read_text()
        assert h.client.get(current+'/edit-layout').status_code == 200
