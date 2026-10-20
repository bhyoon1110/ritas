"""Four factual traceability covers, with original bodies and edit routes retained."""
from copy import deepcopy
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation

from ahn import ppt_report
from ahn.report_branding import NS
from ahn.report_cover import COVER_NAME
from ahn.report_traceability import TRACE_NAME, cover_shapes
from ahn.tests.test_ppt_report import _project_with_sections
from app.report import renderers, tem_editable
from test_posco_report_template import sample_document
from xrd_workbench.test_posco_pptx import fixture
from xrd_workbench.test_sample_slide import IMAGE
from app.xrd_workbench.render.posco_pptx import build_posco_deck, named
from app.xrd_workbench.render.posco_preview import deck_pages


def shape_map(root):
    return {s.find('p:nvSpPr/p:cNvPr', NS).get('name'): s
            for s in root.findall('p:cSld/p:spTree/p:sp', NS)}


@pytest.mark.parametrize('code', ['FTIR', 'RAMAN'])
def test_spectrum_cover_precedes_six_traceability_panels(tmp_path, code):
    document = sample_document(code, tmp_path / 'absent.png')
    root = ET.fromstring(renderers._pptx_title_slide(document))
    shapes = shape_map(root)
    assert not any(name.startswith('Trace ') for name in shapes)
    assert {'Report title','Cover sampleNames','Cover client','Cover requestNumber','Cover analyst','Cover completionDate'} <= set(shapes)
    root = ET.fromstring(renderers._pptx_traceability_slide(document))
    shapes = shape_map(root)
    assert len([name for name in shapes if name.startswith('Trace panel ')])==6
    for key in ('requestNumber','equipmentCode','measuredDate','reviewer','completionDate','recordId'):
        assert 'Trace '+key in shapes
    assert '미기재' in ''.join(shapes['Trace measuredDate'].itertext())
    assert '미기재' in ''.join(shapes['Trace reviewer'].itertext())
    assert '2026-10-15' in ''.join(shapes['Trace completionDate'].itertext())


@pytest.mark.parametrize('template_available', [True,False])
def test_tem_cover_precedes_original_body_and_revision_updates_one_cover(tmp_path, monkeypatch, template_available):
    if not template_available:
        monkeypatch.setattr(ppt_report,'TEMPLATE_PATH',tmp_path/'absent.pptx')
    project = _project_with_sections(tmp_path, {'tem','stem','coating'})
    original = deepcopy(project)
    path = ppt_report.build_pptx(project,tmp_path/'tem.pptx')
    prs = Presentation(path)
    assert len(prs.slides) == 5
    assert prs.slides[0].element.cSld.get('name') == COVER_NAME
    assert prs.slides[1].element.cSld.get('name') == TRACE_NAME
    assert all(s.element.cSld.get('name') not in {COVER_NAME,TRACE_NAME} for s in list(prs.slides)[2:])
    assert any('TEM 이미지 분석결과' in s.text for s in prs.slides[2].shapes if s.has_text_frame)
    model = tem_editable.initial_model(path,{})
    assert model['completionDate'] == '2026-07-08'
    _,layout,_ = tem_editable.inspect_pptx(path)
    refs = {ref for r in layout['pages'][0]['regions'] for ref in r['refs']}
    assert {'title','completionDate','transfer/requestNumber'} <= refs
    assert 'transfer/operatorId' not in refs
    # Repeated review never re-runs OCR or inserts another cover. A blank value
    # remains blank instead of becoming a placeholder or fabricated request.
    for index in range(2):
        model.update(title=f'TEM 검토 완료 {index}', completionDate='2026-10-20')
        model['transfer'].update(requestNumber='REQ-NEW',operatorId='검토자',equipmentCode='')
        revised = tmp_path/f'revision-{index}.pptx'
        tem_editable.revise_pptx(path,revised,model)
        with ZipFile(path) as before, ZipFile(revised) as after:
            for n in before.namelist():
                if n.startswith('ppt/media/') or (n.startswith('ppt/slides/slide') and n.endswith('.xml') and n not in {'ppt/slides/slide1.xml','ppt/slides/slide2.xml'}):
                    left,right = before.read(n),after.read(n)
                    if n.endswith('.xml'):
                        # The existing editor reserializes XML namespace prefixes.
                        left,right = ET.tostring(ET.fromstring(left)),ET.tostring(ET.fromstring(right))
                    assert left==right, n
        prs = Presentation(revised)
        assert len(prs.slides)==5
        content = '\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
        assert model['title'] in content and 'REQ-NEW' in content and '2026-10-20' in content
        assert 'TEM-EDGE-01' not in content
        path = revised
        model = tem_editable.initial_model(path,{})
    assert project == original


def test_xrd_cover_uses_reference_cards_preserves_all_metadata_and_edit_targets():
    project,doc = fixture()
    doc['metadata'].update(client='검증 의뢰기관',author='작성자',analysis_period='2026-10-01 ~ 2026-10-20')
    before=deepcopy((project,doc))
    deck=build_posco_deck(project,doc,lambda _:IMAGE)
    cover=deck.pages[1][0]
    for name in ('Report title','Report eyebrow','Trace notice'):
        assert len(named(cover,name))==1
    assert len([s for s in cover.findall('p:cSld/p:spTree/p:sp',NS)
                if s.find('p:nvSpPr/p:cNvPr',NS).get('name','').startswith('Trace panel ')])==6
    text=''.join(cover.itertext())
    for value in ('REQ-123','XRD-01','QA-OP','검증 의뢰기관','작성자','2026-10-01 ~ 2026-10-20'):
        assert value in text
    assert deck.edits[named(cover,'Report title')[0]]==dict(tab='overview',field='title')
    pages=deck_pages(deck)
    assert len(pages)==5 and pages[0]['kind']=='cover' and pages[1]['kind']=='traceability'
    assert (project,doc)==before
    for spec in cover_shapes('XRD','제목',{},9906000,6858000):
        assert spec['x']+spec['cx']<=9906000
        assert spec['y']+spec['cy']<6630826  # Leave the original RIST logo clear.


def test_old_tem_report_without_cover_remains_editable(tmp_path):
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Inches
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='기존 TEM 결과'
    source=tmp_path/'old.pptx';target=tmp_path/'revised.pptx'
    prs.save(source)
    model=tem_editable.initial_model(source,{})
    model['sections'][0]['paragraphs'][0]='기존 보고서 수정'
    tem_editable.revise_pptx(source,target,model)
    revised=Presentation(target)
    assert len(revised.slides)==1
    assert revised.slides[0].shapes[0].text=='기존 보고서 수정'


def test_long_xrd_cover_fields_fit_without_truncation_or_an_extra_page():
    project,doc=fixture()
    doc['metadata'].update(client='긴'*120,report_number='의'*80,
                           equipment_code='장'*100,operator_id='실'*100)
    deck=build_posco_deck(project,doc,lambda _:IMAGE)
    assert len(deck.pages)==5
    cover=deck.pages[1][0]
    text=''.join(cover.itertext())
    for key in ('client','report_number','equipment_code','operator_id'):
        assert doc['metadata'][key] in text
    for shape in cover.findall('p:cSld/p:spTree/p:sp',NS):
        xf=shape.find('p:spPr/a:xfrm',NS)
        assert int(xf.find('a:off',NS).get('y'))+int(xf.find('a:ext',NS).get('cy'))<6630826


def test_tem_cover_aliases_share_detail_fields_after_repeated_revisions(tmp_path):
    path = ppt_report.build_pptx(_project_with_sections(tmp_path,{'tem'}),tmp_path/'source.pptx')
    for index in range(2):
        model = tem_editable.initial_model(path,{})
        sections,layout,_ = tem_editable.inspect_pptx(path)
        assert layout['pages'][1]['title']=='분석 이력 요약'
        for key,label in (('client','의뢰기관'),('analyst','분석 담당자')):
            section = next(s for s in model['sections'] if s['heading']=='표지 · '+label)
            assert sum(s['heading']=='표지 · '+label for s in sections)==1
            section['paragraphs'][0] = f'{label} 변경 {index}' if index==0 else ''
            ref = f"section/{section['sectionId']}/paragraphs/0"
            for page in layout['pages'][:2]:
                assert sum(ref in r['refs'] for r in page['regions'])==1
        model.update(title=f'검토한 보고서 {index}',completionDate='2026-10-20')
        model['transfer']['requestNumber']='REQ-EDIT'
        target = tmp_path/f'edited-{index}.pptx'
        tem_editable.revise_pptx(path,target,model)
        tem_editable.stamp_record(target,f'new-job-{index}','original-job')
        with ZipFile(target) as z:
            cover=shape_map(ET.fromstring(z.read('ppt/slides/slide1.xml')))
            detail=shape_map(ET.fromstring(z.read('ppt/slides/slide2.xml')))
            for key,label in (('client','의뢰기관'),('analyst','분석 담당자')):
                value=f'{label} 변경 {index}' if index==0 else '미기재'
                assert value in ''.join(cover['Cover '+key].itertext())
                assert value in ''.join(detail['Trace '+key].itertext())
            for shapes in (cover,detail): assert model['title'] in ''.join(shapes['Report title'].itertext())
            assert f'new-job-{index}' in ''.join(detail['Trace recordId'].itertext())
        path=target


def test_legacy_single_trace_cover_stays_single_on_edit(tmp_path):
    path = ppt_report.build_pptx(_project_with_sections(tmp_path,{'tem'}),tmp_path/'new.pptx')
    prs=Presentation(path)
    # Model the previous release: the six-panel page was the only cover.
    slide_id=prs.slides._sldIdLst[0]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides[0].element.cSld.set('name',COVER_NAME)
    old=tmp_path/'legacy.pptx';prs.save(old)
    model=tem_editable.initial_model(old,{})
    model['title']='기존 보고서 수정'
    target=tmp_path/'legacy-revised.pptx'
    tem_editable.revise_pptx(old,target,model)
    tem_editable.stamp_record(target,'legacy-revision')
    reopened=Presentation(target)
    assert len(reopened.slides)==2
    assert tem_editable.initial_model(target,{})['title']==model['title']
    assert 'legacy-revision' in '\n'.join(s.text for s in reopened.slides[0].shapes if s.has_text_frame)


@pytest.mark.parametrize('code',['FTIR','RAMAN'])
def test_spectrum_cover_and_traceability_share_edit_refs_and_values(tmp_path,code):
    from app.report import editable
    from app.report.direct_edit import edit_pages
    from ahn.report_traceability import FIELDS
    doc=sample_document(code,tmp_path/'missing.png')
    editable.ensure_editable_fields(doc,{},code,editable.transfer_values({},'REQ-1','EQ-1','SSO-ID'))
    model=editable.editor_document(doc)
    table=next(s for s in model['sections'] if s['sectionId']=='report_traceability')
    for index,(key,_) in enumerate(FIELDS):
        if key in {'client','analyst'}: table['rows'][index][1]='수정한 '+key
    editable.apply_edits(doc,model,code)
    path=tmp_path/'report.pptx';renderers.render_pptx(doc,path)
    layout=edit_pages(doc,path)
    assert layout['pages'][1]['title']=='분석 이력 요약'
    with ZipFile(path) as z:
        for index,prefix in ((0,'Cover '),(1,'Trace ')):
            shapes=shape_map(ET.fromstring(z.read(f'ppt/slides/slide{index+1}.xml')))
            for row,(key,_) in enumerate(FIELDS):
                if key not in {'client','analyst'}: continue
                assert '수정한 '+key in ''.join(shapes[prefix+key].itertext())
                assert any(f'section/report_traceability/rows/{row}/1' in r['refs'] for r in layout['pages'][index]['regions'])


def test_xrd_cover_and_traceability_share_edit_targets():
    project,doc=fixture()
    doc['metadata'].update(client='동일 의뢰기관',author='동일 작성자')
    deck=build_posco_deck(project,doc,lambda _:IMAGE)
    for key,field in (('client','client'),('analyst','author'),('requestNumber','report_number'),('completionDate','date')):
        for index,prefix in ((0,'Cover '),(1,'Trace ')):
            shape=named(deck.pages[index][0],prefix+key)[0]
            assert doc['metadata'][field] in ''.join(shape.itertext())
            assert deck.edits[shape]==dict(tab='overview',field=field)
