from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from PIL import Image

from app import ahn_web, tem_report_editor
from app.report import tem_canvas, tem_editable, canvas
from test_tem_report_editor import report, editable, wait

HEADERS = {'X-Requested-With': 'RIST-Report-Editor'}
NS = tem_editable.NS


def revise(h, url, model):
    response = h.client.post(url+'/revisions', json=editable(model), headers=HEADERS)
    assert response.status_code == 202, response.text
    state = wait(h, response.json())
    assert state['status'] == 'completed', state
    job = ahn_web._ahn_report_jobs[state['jobId']]
    return '/api/v1/tem/report/jobs/'+job.job_id, job


def test_pages_objects_images_and_restore_survive_repeated_revisions(report):
    h = report
    source_bytes = h.job.pptx_path.read_bytes()
    model = h.client.get(h.url+'/editor-data').json()['document']
    layout = h.client.get(h.url+'/edit-layout').json()
    assert layout['pageEditing'] and layout['canvasEditing']
    ids = [p['id'] for p in layout['pages']]
    assert len(set(ids)) == 4
    original = next(p for p in layout['pages'] if any(o['kind']=='image' for o in p['objects']))
    picture = next(o for o in original['objects'] if o['kind']=='image')
    png = BytesIO(); Image.new('RGB', (80, 40), 'blue').save(png, 'PNG')
    response = h.client.post(h.url+'/editor-images', content=png.getvalue(), headers=HEADERS)
    assert response.status_code == 200, response.text
    asset = response.json()['asset']
    model['pageEdits'] = dict(hidden=[ids[1]], order=[ids[3],ids[0],'added-review',ids[2]], added=[
        dict(id='added-review', after=ids[0], title='TEM 검토 내용', body='추가 검토 결과')])
    model['canvasEdits'] = {original['id']: [dict(id=picture['id'], hidden=True)], 'added-review': [
        dict(id='object-note', kind='text', text='검토 완료', x=.12,y=.65,w=.3,h=.12,bold=True,color='#123456'),
        dict(id='object-line', kind='line', x=.1,y=.8,w=.7,h=.01,dash='dash',stroke='#FF5500'),
        dict(id='object-picture', kind='image', asset=asset,x=.65,y=.3,w=.2,h=.2)]}
    url, job = revise(h, h.url, model)
    actual = h.client.get(url+'/edit-layout').json()
    assert [p['id'] for p in actual['pages']] == model['pageEdits']['order']
    added = next(p for p in actual['pages'] if p['id']=='added-review')
    assert {o['id'] for o in added['objects']} >= {'object-note','object-line','object-picture'}
    assert any(r['refs']==['page/added-review/title'] for r in added['regions'])
    assert not any(o['id']==picture['id'] for p in actual['pages'] for o in p['objects'])
    with ZipFile(job.pptx_path) as ppt:
        from lxml import etree
        # LibreOffice's OPC loader requires the standard default Types namespace.
        assert etree.fromstring(ppt.read('[Content_Types].xml')).prefix is None
        assert len(tem_editable._slides(ppt)) == 4
        assert 'ppt/slides/slide2.xml' not in ppt.namelist()
        assert sum(n.startswith('ppt/media/canvas-') for n in ppt.namelist()) == 1
        assert len(ppt.namelist()) == len(set(ppt.namelist()))
        with ZipFile(h.job.pptx_path) as source:
            for name in source.namelist():
                if name.startswith(('ppt/slideLayouts/', 'ppt/slideMasters/', 'ppt/media/')):
                    assert ppt.read(name) == source.read(name)
    # Current editor state and immutable base survive registry reload/history.
    ahn_web.tem_history.save(job)
    with ahn_web._ahn_report_jobs_lock:
        ahn_web._ahn_report_jobs.pop(job.job_id)
        ahn_web.tem_history._loaded.discard(ahn_web.tem_history.root(h.app.state.settings))
    reopened = h.client.get(url+'/editor-data').json()['document']
    url2, job2 = revise(h, url, reopened)
    assert h.client.get(url2+'/edit-layout').json() == actual
    assert tem_canvas.base_path(job2).read_bytes() == source_bytes
    # Clear absolute patches and restore deleted pages: original picture/geometry returns.
    reopened['canvasEdits'] = {}
    reopened['pageEdits'] = dict(hidden=[], order=ids, added=[])
    restored_url, restored = revise(h, url2, reopened)
    restored_layout = h.client.get(restored_url+'/edit-layout').json()
    assert [p['id'] for p in restored_layout['pages']] == ids
    assert any(o['id']==picture['id'] for p in restored_layout['pages'] for o in p['objects'])
    assert h.job.pptx_path.read_bytes() == source_bytes
    assert not (restored.work_dir/'input').exists()  # Never copy multi-GB raw files.
    with ZipFile(restored.package_path) as package:
        assert not any('base.pptx' in n or 'editor-assets' in n for n in package.namelist())
        assert package.read('tem-report.pptx') == restored.pptx_path.read_bytes()


def test_native_table_edits_and_actual_tem_page_dimensions(report):
    from pptx import Presentation
    from pptx.util import Inches
    h = report
    base = tem_canvas.base_path(h.job)
    prs = Presentation(base)
    table = prs.slides[-1].shapes.add_table(2,2,Inches(.5), Inches(1), Inches(3), Inches(1)).table
    table.cell(0,0).text = '항목'; table.cell(0,1).text = '결과'
    table.cell(1,0).text = '검토'; table.cell(1,1).text = '전'
    prs.save(h.job.pptx_path)
    base.unlink()  # Fixture only: regenerate the baseline containing a native table.
    tem_report_editor.prepare_review(h.job, h.job.manifest)
    model = h.client.get(h.url+'/editor-data').json()['document']
    layout = h.client.get(h.url+'/edit-layout').json()
    page = layout['pages'][-1]
    obj = next(o for o in page['objects'] if o['kind']=='table')
    model['canvasEdits'] = {page['id']:[dict(id=obj['id'],x=.2,y=.4,w=.55,h=.2,
        tableCells=[dict(row=1,col=1,text='검토 후',bold=True,fill='#EEEEFF')])]}
    url, job = revise(h,h.url,model)
    result = next(o for o in h.client.get(url+'/edit-layout').json()['pages'][-1]['objects'] if o['id']==obj['id'])
    for key, value in dict(x=.2,y=.4,w=.55,h=.2).items():
        assert result[key] == pytest.approx(value, abs=1e-6)
    assert next(c for c in result['cells'] if c['row']==1 and c['col']==1)['text']=='검토 후'
    assert next(c for c in result['cells'] if c['row']==1 and c['col']==1)['bold']
    assert len(Presentation(job.pptx_path).slides[-1].shapes)==len(prs.slides[-1].shapes)
    cleared=h.client.get(url+'/editor-data').json()['document']
    cleared['canvasEdits'][page['id']][0]['tableCells']=[dict(row=r,col=c,text='') for r in range(2) for c in range(2)]
    cleared_url,_=revise(h,url,cleared)
    objects=h.client.get(cleared_url+'/edit-layout').json()['pages'][-1]['objects']
    assert any(o['id']==obj['id'] and o['kind']=='table' for o in objects)


@pytest.mark.parametrize('attack', ['unknown-page','duplicate-order','empty-deck','unknown-object','locked-record',
                                  'missing-asset','bad-anchor','too-many-objects'])
def test_structural_edits_rejected_before_queueing(report, monkeypatch, attack):
    h = report
    model = h.client.get(h.url+'/editor-data').json()['document']
    layout = h.client.get(h.url+'/edit-layout').json()
    ids = [p['id'] for p in layout['pages']]
    if attack=='unknown-page': model['pageEdits']['hidden']=['nonexistent']
    if attack=='duplicate-order': model['pageEdits']['order']=[ids[0],ids[0]]
    if attack=='empty-deck': model['pageEdits']['hidden']=ids
    if attack=='unknown-object': model['canvasEdits']={ids[0]:[dict(id='shape-'+'a'*20,hidden=True)]}
    if attack=='locked-record':
        with ZipFile(h.job.pptx_path) as archive:
            payload=tem_canvas._read(archive)[1]
            payload.pop('lockedShapeIds')
            obj=next(o for _,o in canvas.objects(payload)[1] if o['name']=='Trace recordId')
        model['canvasEdits']={ids[1]:[dict(id=obj['id'],text='fake-id')]}
    if attack=='missing-asset': model['canvasEdits']={ids[0]:[dict(id='object-image',kind='image',asset='a'*64)]}
    if attack=='bad-anchor': model['pageEdits']['added']=[dict(id='added-bad',after='missing',title='',body='')]
    if attack=='too-many-objects': model['canvasEdits']={ids[0]:[dict(id='object-'+str(i),kind='text') for i in range(101)]}
    def unexpected(*args, **kwargs): pytest.fail('invalid edit must not queue work')
    monkeypatch.setattr(ahn_web._ahn_report_executor, 'submit', unexpected)
    response = h.client.post(h.url+'/revisions', json=editable(model), headers=HEADERS)
    assert response.status_code == 422, response.text


def test_image_authorization_origin_size_format_and_thumbnail(report):
    from dataclasses import replace
    from app.auth import AuthContext
    from datetime import datetime, timedelta
    h=report
    assert h.client.post(h.url+'/editor-images', content=b'no').status_code==403
    assert h.client.post(h.url+'/editor-images', content=b'no', headers=HEADERS|{'Origin':'https://foreign.invalid'}).status_code==403
    assert h.client.post(h.url+'/editor-images', content=b'no', headers=HEADERS).status_code==422
    assert h.client.post(h.url+'/editor-images', content=b'x'*(2*1024*1024+1), headers=HEADERS).status_code==413
    thumbnail=h.client.get(h.url+'/edit-slides/0?thumbnail=true')
    with Image.open(BytesIO(thumbnail.content)) as image:
        assert image.width <= 240
    h.app.state.settings=replace(h.app.state.settings, auth_enabled=True)
    h.app.state.actor=AuthContext(user_id='other',login_id='other',display_name='other',email=None,status='ACTIVE',
        session_id='test',session_expires_at=datetime.now()+timedelta(hours=1),sso_authenticated_at=None,
        projects=frozenset({'TEM'}),roles=frozenset(),sso_identity=None)
    assert h.client.post(h.url+'/editor-images', content=b'no', headers=HEADERS).status_code==403


def test_pre_upgrade_job_without_base_keeps_text_and_can_restore(report):
    h=report
    tem_canvas.base_path(h.job).unlink()  # Simulate a persisted pre-upgrade job.
    model=h.client.get(h.url+'/editor-data').json()['document']
    model.pop('pageEdits');model.pop('canvasEdits')
    model['title']='이전 보고서 수정'
    url,job=revise(h,h.url,model)
    assert tem_canvas.base_path(job).is_file()
    assert h.client.get(url+'/edit-layout').json()['pageEditing']
