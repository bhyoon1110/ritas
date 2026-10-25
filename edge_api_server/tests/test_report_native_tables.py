"""Actual PPT table structure, old revisions and author-only editing contracts."""
from copy import deepcopy
import hashlib
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation
import pytest

from app.report import canvas, renderers, pptx_tables
from app.report.direct_edit import edit_pages
from app.report.editable import ensure_editable_fields, transfer_values
from app.report.model import ReportDocument, ReportSection, ReportTable


def document(code='FT-IR'):
    doc=ReportDocument('native-table','표 편집 검증',code,{},'2026-10-26')
    doc.sections=[ReportSection('sample_peaks','시료별 피크',table=ReportTable(
        columns=['시료','Position (cm⁻¹)','Intensity (a.u.)'],
        rows=[['Sample A','1350','1.25'],['Sample A','1580','0.75']],
        merge_columns=[0],column_widths=[.4,.3,.3],alignments=['left','right','right']))]
    return doc


def table_payload(doc):
    return next(p for p in renderers._pptx_slide_payloads(doc) if p.get('editing',{}).get('sectionId')=='sample_peaks')


@pytest.mark.parametrize('code',['FT-IR','RAMAN'])
def test_export_has_one_native_table_with_editable_cells_merges_and_branding(tmp_path,code):
    doc=document(code);target=tmp_path/'report.pptx';renderers.render_pptx(doc,target)
    prs=Presentation(target)
    native=[s for slide in prs.slides for s in slide.shapes if s.has_table]
    assert len(native)==1
    table=native[0].table
    assert len(table.rows)==3 and len(table.columns)==3
    assert table.cell(0,0).text=='시료'
    assert table.cell(1,1).text=='1350' and table.cell(2,2).text=='0.75'
    assert table.cell(1,0).is_merge_origin and table.cell(1,0).span_height==2
    assert table.cell(2,0).is_spanned
    assert str(table.cell(0,0).fill.fore_color.rgb)==renderers.PPTX_NAVY
    assert table.cell(0,0).text_frame.paragraphs[0].runs[0].font.bold
    assert not any(s.name.startswith(('Cell ','Header ')) for slide in prs.slides for s in slide.shapes)
    # Native table remains native after a standard PPT object-model save.
    table.cell(1,1).text='1351.25';prs.save(tmp_path/'edited.pptx')
    assert [s for slide in Presentation(tmp_path/'edited.pptx').slides for s in slide.shapes if s.has_table][0].table.cell(1,1).text=='1351.25'
    with ZipFile(target) as z:
        assert 'ppt/media/posco-banner.jpeg' in z.namelist() and 'ppt/media/rist-logo.png' in z.namelist()


def test_native_table_transform_cell_edits_regions_and_undo(tmp_path):
    doc=document();payload=table_payload(doc)
    oid=next(i['id'] for _,i in canvas.objects(payload)[1] if i['kind']=='table')
    original=deepcopy(doc.to_dict())
    doc.canvas_edits={payload['pageId']:[dict(id=oid,x=.07,y=.25,w=.8,h=.36,
        tableCells=[dict(row=1,col=1,text='1352\n확인 <&>',color='#123456',bold=True,fontSize=13,
                        fill='#FFFFDD',stroke='#CC6633',lineWidth=2,dash='dash',align='center')])]}
    target=tmp_path/'report.pptx';renderers.render_pptx(doc,target)
    page=next(p for p in edit_pages(doc,target)['pages'] if p['id']==payload['pageId'])
    table=next(i for i in page['objects'] if i['kind']=='table')
    assert (table['x'],table['y'],table['w'],table['h'])==pytest.approx((.07,.25,.8,.36))
    cell=next(c for c in table['cells'] if (c['row'],c['col'])==(1,1))
    assert cell['text']=='1352\n확인 <&>' and cell['color']=='#123456' and cell['align']=='center'
    assert cell['bold'] and cell['fill']=='#FFFFDD' and cell['dash']=='dash' and cell['lineWidth']==2
    merged=next(r for r in page['regions'] if 'section/sample_peaks/rows/0/0' in r['refs'])
    assert merged['refs']==['section/sample_peaks/rows/0/0','section/sample_peaks/rows/1/0']
    assert merged['x']==pytest.approx(.07) and merged['w']==pytest.approx(.32)
    assert merged['h']==pytest.approx(.24)
    assert doc.to_dict()['sections']==original['sections']  # Free edits are presentation-only.
    ppt=Presentation(target)
    table=next(s.table for slide in ppt.slides for s in slide.shapes if s.has_table)
    assert table.cell(1,1).text=='1352\n확인 <&>'
    assert sum(r.height for r in table.rows)==round(.36*renderers.PPTX_SLIDE_H)
    doc.canvas_edits={};renderers.render_pptx(doc,target)
    assert next(s.table for slide in Presentation(target).slides for s in slide.shapes if s.has_table).cell(1,1).text=='1350'


def test_legacy_per_cell_edits_survive_with_explicit_reflow_notice(tmp_path):
    doc=document();payload=table_payload(doc)
    oid='shape-'+hashlib.sha256(b'Cell 0-1:0').hexdigest()[:20]
    native=next(i['id'] for _,i in canvas.objects(payload)[1] if i['kind']=='table')
    # Native edits win even if old cell patches are ordered later.
    doc.canvas_edits={payload['pageId']:[dict(id=native,tableCells=[dict(row=1,col=1,color='#001122')]),
                                       dict(id=oid,text='기존 편집',color='#112233',x=.2,bold=True)]}
    saved=deepcopy(doc.canvas_edits)
    canvas.apply(doc,[payload],validate_only=True,previous_edits=saved)
    target=tmp_path/'report.pptx';renderers.render_pptx(doc,target)
    page=next(p for p in edit_pages(doc,target)['pages'] if p['id']==payload['pageId'])
    assert page['legacyTableLayout']
    table=next(i for i in page['objects'] if i['kind']=='table')
    cell=next(c for c in table['cells'] if (c['row'],c['col'])==(1,1))
    assert cell['text']=='기존 편집' and cell['bold'] and cell['color']=='#001122'
    assert doc.canvas_edits==saved


def test_preexisting_shape_report_uses_its_archived_geometry_until_regenerated(tmp_path):
    doc=document();target=tmp_path/'new.pptx';renderers.render_pptx(doc,target)
    payloads=renderers._pptx_slide_payloads(doc)
    index=next(i for i,p in enumerate(payloads) if p.get('editing',{}).get('sectionId')=='sample_peaks')
    name=f'ppt/slides/slide{index+1}.xml'
    with ZipFile(target) as z:
        files={n:z.read(n) for n in z.namelist()}
    root=ET.fromstring(files[name]);parent=root.find('p:cSld/p:spTree',canvas.NS)
    parent.remove(parent.find('p:graphicFrame',canvas.NS))
    old=renderers._pptx_text_shape(10,'Cell 0-1',x=1000000,y=2000000,cx=3000000,cy=500000,lines=['과거 보고서 셀'])
    parent.append(ET.fromstring('<root xmlns:p="'+canvas.NS['p']+'" xmlns:a="'+canvas.NS['a']+'">'+old+'</root>')[0])
    files[name]=ET.tostring(root)
    legacy=tmp_path/'legacy.pptx'
    with ZipFile(legacy,'w') as z:
        for name,data in files.items():z.writestr(name,data)
    page=edit_pages(doc,legacy)['pages'][index]
    assert not any(o['kind']=='table' for o in page['objects'])
    assert any(o['text']=='과거 보고서 셀' for o in page['objects'])
    region=next(r for r in page['regions'] if r['refs']==['section/sample_peaks/rows/0/1'])
    assert region['x']==1000000/renderers.PPTX_SLIDE_W


@pytest.mark.parametrize('cells',[[dict(row=99,col=0,text='x')],
                                 [dict(row=2,col=0,text='merged continuation')],
                                 [dict(row=1,col=1,text='a'),dict(row=1,col=1,text='b')]])
def test_invalid_native_cells_fail_closed(cells):
    doc=document();payload=table_payload(doc)
    oid=next(i['id'] for _,i in canvas.objects(payload)[1] if i['kind']=='table')
    doc.canvas_edits={payload['pageId']:[dict(id=oid,tableCells=cells)]}
    with pytest.raises(ValueError):canvas.apply(doc,[payload],validate_only=True)


@pytest.mark.parametrize('cell',[dict(row=-1,col=0),dict(row=0,col=0,text='x'*4001),
                                dict(row=0,col=0,fontSize=float('nan')),dict(row=0,col=0,asset='../file')])
def test_nested_cell_schema_is_bounded(cell):
    with pytest.raises(ValueError):canvas.ObjectEdit(id='object-test',tableCells=[cell])


def test_report_revision_does_not_persist_null_cell_style_overrides():
    from app.report.editable import editor_document
    from app.report_editor import ReportEdit, _validate_changes
    from tests.test_report_editor import edit_request
    doc=document();ensure_editable_fields(doc,{},'FT-IR',transfer_values({}))
    payload=table_payload(doc)
    oid=next(i['id'] for _,i in canvas.objects(payload)[1] if i['kind']=='table')
    request=edit_request(editor_document(doc))
    request['canvasEdits']={payload['pageId']:[dict(id=oid,tableCells=[dict(row=1,col=1,text='수정')])]}
    values=_validate_changes(ReportEdit(**request),doc)
    assert values['canvasEdits']==request['canvasEdits']


def test_all_paginated_cells_remain_native_without_losing_last_values(tmp_path):
    doc=document('RAMAN');doc.sections[0].section_id='raman_peaks'
    doc.sections[0].table.rows=[['시료 '+str(i),'1350.125','강도 확인 문구 '*20+str(i)] for i in range(25)]
    ensure_editable_fields(doc,{},'RAMAN',transfer_values({}))
    target=tmp_path/'many.pptx';renderers.render_pptx(doc,target)
    tables=[s.table for slide in Presentation(target).slides for s in slide.shapes if s.has_table]
    assert len(tables)>5
    texts=[cell.text for table in tables for row in table.rows for cell in row.cells]
    assert any('시료 24' in text for text in texts) and any(text.endswith('24') for text in texts)
    assert any('section/raman_peaks/rows/24/2' in r['refs'] for p in edit_pages(doc,target)['pages'] for r in p['regions'])


def test_resized_grid_dimensions_and_cell_border_xml_order():
    doc=document();payload=table_payload(doc)
    frame,item=next((s,i) for s,i in canvas.objects(payload)[1] if i['kind']=='table')
    cell=next(c for row,col,c,_ in pptx_tables.cells(frame) if (row,col)==(1,1))
    canvas._edit_cell(cell,dict(fill='#FFFFCC',stroke='#112233',text='검증',color='#224466'))
    tcpr=cell.find('a:tcPr',canvas.NS)
    assert [n.tag.split('}')[-1] for n in tcpr]==['lnL','lnR','lnT','lnB','solidFill']
    pptx_tables.resize(frame,1000001,777777)
    table=frame.find(pptx_tables.TABLE_PATH,canvas.NS)
    assert sum(int(n.get('w')) for n in table.findall('a:tblGrid/a:gridCol',canvas.NS))==1000001
    assert sum(int(n.get('h')) for n in table.findall('a:tr',canvas.NS))==777777
