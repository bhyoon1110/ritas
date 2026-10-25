"""Isolated loopback fixture: real TEM PPT, LibreOffice PDF, revision and ZIP."""
import json
import os
from pathlib import Path
import socket
import sys

root = Path(sys.argv[1])
os.environ['RIST_STORAGE_ROOT'] = str(root/'storage')
os.environ['RIST_ERROR_ARCHIVE_ROOT'] = str(root/'errors')
os.environ['RIST_USAGE_LOG_ROOT'] = str(root/'usage')
os.environ['RIST_AUTH_ENABLED'] = 'false'

from app import ahn_web, tem_report_editor
from app.storage import atomic_write_json
from ahn.ppt_report import build_pptx
from ahn.tests.test_ppt_report import _project_with_sections
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw
import uvicorn

app = ahn_web.create_tem_preview_app()
workspace = ahn_web.tem_history.new_workspace(app.state.settings)
inputs = workspace/'input'; inputs.mkdir(parents=True)
job = ahn_web._create_ahn_job(inputs, workspace, None, app.state.settings, owner_user_id='test-author')
job.output_dir.mkdir()
project = _project_with_sections(inputs, {'tem', 'coating'})
build_pptx(project, job.pptx_path)
prs = Presentation(job.pptx_path)
slide = prs.slides.add_slide(prs.slides[0].slide_layout)
shape = slide.shapes.add_table(2,2,Inches(.6),Inches(1.2),Inches(5),Inches(1.1))
for row, values in zip(shape.table.rows, [['검토 항목','결과'],['표 편집','이전 값']]):
    for cell,value in zip(row.cells,values): cell.text = value
prs.save(job.pptx_path)
atomic_write_json(job.analysis_path, project)
manifest = {'summary': {'temSampleCount':1,'coatingSampleCount':1},'pptx':'tem-report.pptx'}
atomic_write_json(job.manifest_path, manifest)
tem_report_editor.prepare_review(job, manifest)
if job.review_error: raise RuntimeError(job.review_error)
with Image.new('RGB', (240,120), '#1379A8') as image:
    ImageDraw.Draw(image).text((24,48), 'TEM review image', fill='white')
    image.save(root/'test-image.png')
ahn_web._build_package(job.output_dir, job.package_path)
ahn_web._set_job_state(job, status='completed', manifest=manifest)


@app.get('/api/v1/auth/me')
def me():
    return {'userId':'test-author','projects':['TEM'],'roles':[]}


@app.get('/api/v1/requests')
def requests():
    return {'items':[]}


sock = socket.socket(); sock.bind(('127.0.0.1',0))
print(json.dumps({'port':sock.getsockname()[1], 'jobId':job.job_id}), flush=True)
uvicorn.Server(uvicorn.Config(app,log_level='warning')).run(sockets=[sock])
