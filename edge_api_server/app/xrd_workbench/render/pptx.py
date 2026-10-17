"""Editable 16:9 PPTX from the shared page model: one text box per pre-wrapped line, source images
as separate pictures, relative source links, and the same title band and footer as the PDF (they
are page items like any other)."""

import io
from datetime import datetime, timezone

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..config import load_settings
from ..render.fonts import family_of, font_names, pptx_family, segments
from ..render.items import link_target
from ..render.theme import PAGE_H, PAGE_W

EMU_PER_MM = 36000
_FROM_SETTINGS = object()  # render_pptx: use XRD_PPTX_TEMPLATE when none is passed


def _mm(value: float) -> Emu:
    return Emu(round(value * EMU_PER_MM))


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip('#').upper())


def _set_typeface(run, family: str) -> None:
    """Latin and East Asian typeface; without <a:ea> Korean glyphs fall back to the theme font."""
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea'):
        node = rpr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rpr, qn(tag))
        node.set('typeface', family)


def _text(
    slide,
    content,
    *,
    x,
    y,
    w,
    h,
    size,
    color,
    family,
    bold=False,
    leading=None,
    align=PP_ALIGN.LEFT,
    name,
    href=None,
    scale=(1, 0, 0),
):
    """One absolutely positioned, editable text box per line (no wrapping or auto-fit). `scale`
    maps the page model's millimetres onto a template slide of another size."""
    s, ox, oy = scale
    size = size * s
    lines = str(content).split('\n')
    step = (leading if leading is not None else h / len(lines)) * s
    for index, line in enumerate(lines):
        height = h * s if len(lines) == 1 else step
        box = slide.shapes.add_textbox(
            _mm(x * s + ox), _mm(y * s + oy + index * step), _mm(w * s), _mm(height)
        )
        box.name = f'{name} line {index + 1}'
        frame = box.text_frame
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        primary = font_names()[1 if bold else 0]
        # One run per font: characters the configured font lacks declare the bundled fallback
        # family (Noto Sans / Noto Sans Math) instead of relying on the viewer's substitution.
        for font_name, part in segments(line, bold) or [(primary, line)]:
            run = paragraph.add_run()
            run.text = part
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = _rgb(color)
            _set_typeface(run, family if font_name == primary else family_of(font_name))
            if href:
                run.hyperlink.address = href


def _rect(slide, item, name, scale=(1, 0, 0)):
    s, ox, oy = scale
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        _mm(item['x'] * s + ox),
        _mm(item['y'] * s + oy),
        _mm(item['w'] * s),
        _mm(item['h'] * s),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(item['fill'])
    shape.line.fill.background()
    shape.shadow.inherit = False


def _open_deck(template, layout_name):
    """(deck, layout, scale): python-pptx's blank 16:9 deck, or the company template with its own
    slides removed. scale = (factor, left, top): the page model's millimetres map onto the
    template's slide with one factor (proportions stay) and the content centred."""
    if template is None:
        deck = Presentation()
        deck.slide_width, deck.slide_height = _mm(PAGE_W), _mm(PAGE_H)
        # The default template is a 4:3 on-screen deck made by python-pptx; this is a 16:9 report.
        deck.part._element.sldSz.attrib.pop('type', None)
        return deck, deck.slide_layouts[6], (1, 0, 0)
    try:
        deck = Presentation(str(template))
    except Exception as exc:
        raise ValueError(
            f'PPT 양식을 열지 못했습니다: {template} (있는 .pptx 파일인지 확인하세요. '
            '.potx는 PowerPoint에서 .pptx로 저장해서 쓰세요).'
        ) from exc
    slide_ids = deck.slides._sldIdLst
    for slide_id in list(slide_ids):
        deck.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    layouts = list(deck.slide_layouts)
    if not layouts:
        raise ValueError(f'PPT 양식에 슬라이드 레이아웃이 없습니다: {template}')
    if layout_name:
        layout = next((x for x in layouts if x.name == layout_name), None)
        if layout is None:
            raise ValueError(
                f'PPT 양식에 "{layout_name}" 레이아웃이 없습니다. 있는 레이아웃: '
                + ', '.join(x.name for x in layouts)
            )
    else:
        # The emptiest layout: the report places its own text and pictures.
        layout = min(layouts, key=lambda x: len(x.placeholders))
    factor = min(deck.slide_width / _mm(PAGE_W), deck.slide_height / _mm(PAGE_H))
    left = (deck.slide_width - _mm(PAGE_W) * factor) / 2 / EMU_PER_MM
    top = (deck.slide_height - _mm(PAGE_H) * factor) / 2 / EMU_PER_MM
    return deck, layout, (factor, left, top)


def render_pptx(
    pages: list[dict],
    document: dict,
    font_family: str | None = None,
    link_base: str | None = None,
    storage_id: str | None = None,
    template=_FROM_SETTINGS,
    template_layout: str | None = None,
) -> bytes:
    """Render the shared page model as an editable 16:9 deck (same geometry as the PDF)."""
    family = font_family or pptx_family()
    settings = load_settings()
    if template is _FROM_SETTINGS:
        template, template_layout = settings.pptx_template, settings.pptx_layout
    deck, blank, scale = _open_deck(template, template_layout)
    # White title text suits the temporary navy band; a template's own band may need another colour.
    title_color = settings.pptx_title_color if template is not None else None
    props = deck.core_properties
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    props.title = document['metadata']['title']
    props.author = document['metadata'].get('author') or ''
    props.last_modified_by = props.author
    props.comments = ''
    props.created = props.modified = now
    for page in pages:
        slide = deck.slides.add_slide(blank)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        for i, item in enumerate(page['items']):
            if item['type'] == 'rect':
                if template is not None and item.get('role') == 'header':
                    continue  # the company template draws its own title band
                _rect(
                    slide,
                    item,
                    f'Band {i}' if item.get('role') == 'header' else f'Rectangle {i}',
                    scale,
                )
            elif item['type'] == 'text':
                _text(
                    slide,
                    '\n'.join(item['lines']),
                    x=item['x'],
                    y=item['y'],
                    w=item['w'],
                    h=item['h'],
                    size=item['size'],
                    bold=item['bold'],
                    color=(title_color or item['color'])
                    if item.get('role') == 'header'
                    else item['color'],
                    leading=item['leading'],
                    family=family,
                    align=PP_ALIGN.RIGHT if item.get('align') == 'right' else PP_ALIGN.LEFT,
                    name=f'Text {i}',
                    scale=scale,
                )
            elif item['type'] == 'image':
                picture = slide.shapes.add_picture(
                    io.BytesIO(item['data']),
                    _mm(item['x'] * scale[0] + scale[1]),
                    _mm(item['y'] * scale[0] + scale[2]),
                    _mm(item['w'] * scale[0]),
                    _mm(item['h'] * scale[0]),
                )
                picture.name = f'Image {i}'
                picture._element.nvPicPr.cNvPr.set('descr', item['filename'])
            elif item['type'] == 'link':
                _text(
                    slide,
                    item['label'] + ' ↗',
                    x=item['x'],
                    y=item['y'],
                    w=item['w'],
                    h=item['h'],
                    size=item['size'],
                    color='#367ab1',
                    family=family,
                    leading=item.get('leading'),
                    name=f'Source link {i}',
                    href=link_target(item, link_base, storage_id),
                    scale=scale,
                )
        evidence = [
            f'{i["filename"]}: {link_target(i, link_base, storage_id)}'
            for i in page['items']
            if i['type'] == 'link'
        ]
        slide.notes_slide.notes_text_frame.text = '\n'.join(evidence)
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()
