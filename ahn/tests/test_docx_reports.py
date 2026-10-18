from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation

from ahn.analysis import collect_project
from ahn.docx_extract import extract_docx
from ahn.ppt_report import build_pptx


def write_docx(path: Path, color=(220, 40, 40), *, vml=False, absolute=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    image = BytesIO()
    Image.new('RGB', (80, 60), color).save(image, format='PNG')
    drawing = '<v:imagedata r:id="image"/>' if vml else '<a:blip r:embed="image"/>'
    target = '/word/media/image.png' if absolute else 'media/image.png'
    with ZipFile(path, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml',
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:body>' + drawing + '</w:body></w:document>')
        archive.writestr('word/_rels/document.xml.rels',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="image" Target="{target}"/></Relationships>')
        archive.writestr('word/media/image.png', image.getvalue())


def test_collects_loose_and_nested_docx_without_duplicate_or_office_lock_files(tmp_path):
    names = ['loose MAP.docx', 'report/first MAP.docx', 'reports/sample/second MAP.DOCX']
    for name in names + ['reports/~$locked.docx', '.hidden/secret.docx', '__MACOSX/ignored.docx']:
        write_docx(tmp_path / name)
    project = collect_project(tmp_path)
    assert project.summary['edsReportCount'] == 3
    assert [report.path for report in project.eds_reports] == names


@pytest.mark.parametrize('vml,absolute', [(False, False), (True, False), (True, True)])
def test_extracts_instrument_docx_image_formats(tmp_path, vml, absolute):
    source = tmp_path / 'report.docx'
    write_docx(source, vml=vml, absolute=absolute)
    extracted = extract_docx(source, tmp_path / 'extracted')
    assert len(extracted.media_paths) == 1
    assert Image.open(extracted.media_paths[0]).getpixel((0, 0)) == (220, 40, 40)


def test_same_named_documents_have_distinct_slides_and_image_cache(tmp_path):
    write_docx(tmp_path / 'report/sampleA/MAP.docx', (220, 40, 40))
    write_docx(tmp_path / 'report/sampleB/MAP.docx', (40, 40, 220))
    project = collect_project(tmp_path)
    output = build_pptx(project, tmp_path / 'report.pptx')
    slides = Presentation(output).slides
    assert len(slides) == 2
    pixels = []
    for slide in slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        pixels.append(Image.open(BytesIO(pictures[0].image.blob)).convert('RGB').getpixel((10, 10)))
    assert pixels[0][0] > pixels[0][2] + 100
    assert pixels[1][2] > pixels[1][0] + 100


def test_empty_docx_is_not_silently_omitted_from_successful_report(tmp_path):
    source = tmp_path / 'empty.docx'
    with ZipFile(source, 'w') as archive:
        archive.writestr('word/document.xml', '<document/>')
    with pytest.raises(ValueError, match='empty.docx.*이미지 또는 표'):
        build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')


def test_table_only_docx_keeps_all_tables_on_continuation_slides(tmp_path):
    source = tmp_path / 'tables.docx'
    tables = ''.join('<w:tbl><w:tr><w:tc><w:p><w:r><w:t>' + text +
                     '</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
                     for text in ['first', 'second', 'third'])
    with ZipFile(source, 'w') as archive:
        archive.writestr('word/document.xml',
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body>' + tables + '</w:body></w:document>')
    output = build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')
    slides = Presentation(output).slides
    assert len(slides) == 2
    values = [cell.text for slide in slides for shape in slide.shapes if shape.has_table
              for row in shape.table.rows for cell in row.cells]
    assert values == ['first', 'second', 'third']


def test_docx_deduplicates_drawing_vml_fallback_and_ignores_external_images(tmp_path):
    source = tmp_path / 'fallback.docx'
    write_docx(source)
    with ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members['word/document.xml'] = members['word/document.xml'].replace(
        b'</w:body>', b'<v:imagedata r:id="image"/><a:blip r:embed="remote"/></w:body>')
    members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
        b'</Relationships>', b'<Relationship Id="remote" Target="https://example.invalid/private.png" '
        b'TargetMode="External"/></Relationships>')
    with ZipFile(source, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    assert len(extract_docx(source, tmp_path / 'out').media_paths) == 1


def rewrite_parts(source, transform):
    with ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members = transform(members)
    with ZipFile(source, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)


@pytest.mark.parametrize('images', [True, False])
def test_strict_docx_images_and_tables_are_not_treated_as_empty(tmp_path, images):
    source = tmp_path / 'strict MAP.docx'
    write_docx(source)
    def transform(members):
        xml = members['word/document.xml']
        if not images:
            xml = xml.replace(b'<a:blip r:embed="image"/>', b'')
        xml = xml.replace(b'</w:body>', b'<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Oxygen 45%</w:t>'
                          b'</w:r></w:p></w:tc></w:tr></w:tbl></w:body>')
        for old, new in [
            (b'http://schemas.openxmlformats.org/wordprocessingml/2006/main', b'http://purl.oclc.org/ooxml/wordprocessingml/main'),
            (b'http://schemas.openxmlformats.org/drawingml/2006/main', b'http://purl.oclc.org/ooxml/drawingml/main'),
            (b'http://schemas.openxmlformats.org/officeDocument/2006/relationships', b'http://purl.oclc.org/ooxml/officeDocument/relationships'),
        ]:
            xml = xml.replace(old, new)
        members['word/document.xml'] = xml
        return members
    rewrite_parts(source, transform)
    before = source.read_bytes()
    result = extract_docx(source, tmp_path / 'out')
    assert len(result.media_paths) == int(images)
    assert result.tables == [[['Oxygen 45%']]]
    assert not result.diagnostics
    output = build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')
    slides = Presentation(output).slides
    assert len(slides) == 1
    if images:
        assert any(shape.shape_type == 13 for shape in slides[0].shapes)
    else:
        assert any(shape.has_table for shape in slides[0].shapes)
    assert source.read_bytes() == before


def test_vml_office_relid_and_encoded_media_target(tmp_path):
    source = tmp_path / 'legacy MAP.docx'
    write_docx(source, vml=True)
    def transform(members):
        members['word/document.xml'] = members['word/document.xml'].replace(
            b'<v:imagedata r:id="image"/>',
            b'<v:imagedata xmlns:o="urn:schemas-microsoft-com:office:office" o:relid="image"/>')
        members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
            b'media/image.png', b'media/EDS%20map.png')
        members['word/media/EDS map.png'] = members.pop('word/media/image.png')
        return members
    rewrite_parts(source, transform)
    result = extract_docx(source, tmp_path / 'out')
    assert len(result.media_paths) == 1
    assert Image.open(result.media_paths[0]).getpixel((0, 0)) == (220, 40, 40)


@pytest.mark.parametrize('target', ['/media/map.png', '../media/map.png', '/media/EDS%20map.png'])
def test_package_root_media_is_extracted_from_document_relationships(tmp_path, target):
    """Same package layout as the instrument's LTO201 Map1 report."""
    source = tmp_path / 'root MAP.docx'
    write_docx(source)
    def transform(members):
        members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
            b'media/image.png', target.encode())
        name = 'EDS map.png' if '%20' in target else 'map.png'
        members['media/' + name] = members.pop('word/media/image.png')
        # A footer/logo at word/media with the same basename is not the raw map.
        logo = BytesIO(); Image.new('RGB', (20, 20), (0, 0, 220)).save(logo, format='PNG')
        members['word/media/' + name] = logo.getvalue()
        members['media/unreferenced.png'] = logo.getvalue()
        return members
    rewrite_parts(source, transform)
    result = extract_docx(source, tmp_path / 'out')
    assert len(result.media_paths) == 1
    assert not result.diagnostics
    assert Image.open(result.media_paths[0]).getpixel((0, 0)) == (220, 40, 40)
    output = build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')
    assert len(Presentation(output).slides) == 1


@pytest.mark.parametrize('target,external', [
    ('/media/missing.png', False), ('/media/image.png', True),
    ('/media/%2e%2e/secret.png', False), ('../media/../../secret.png', False),
])
def test_root_media_does_not_fetch_external_or_unrelated_or_missing_images(tmp_path, target, external):
    source = tmp_path / 'unavailable.docx'
    write_docx(source)
    def transform(members):
        members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
            b'Target="media/image.png"', f'Target="{target}"'.encode() + (b' TargetMode="External"' if external else b''))
        members['media/image.png'] = members['word/media/image.png']
        members['secret.png'] = members['word/media/image.png']
        return members
    rewrite_parts(source, transform)
    result = extract_docx(source, tmp_path / 'out')
    assert not result.media_paths
    assert '내장 미디어 2개' in result.diagnostics[0]
    with pytest.raises(ValueError, match='이미지 또는 표'):
        build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')


@pytest.mark.parametrize('kind', ['external', 'missing', 'outside-media', 'altChunk'])
def test_unusable_content_has_diagnostics_and_does_not_fall_back_to_orphan_images(tmp_path, kind):
    source = tmp_path / (kind + '.docx')
    write_docx(source)
    def transform(members):
        if kind == 'external':
            members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
                b'Target="media/image.png"', b'Target="https://example.invalid/image.png" TargetMode="External"')
        elif kind == 'missing':
            members.pop('word/media/image.png')
        elif kind == 'outside-media':
            members['word/_rels/document.xml.rels'] = members['word/_rels/document.xml.rels'].replace(
                b'media/image.png', b'media/%2e%2e/secret.png')
            members['word/secret.png'] = members['word/media/image.png']
        else:
            members['word/document.xml'] = members['word/document.xml'].replace(
                b'<a:blip r:embed="image"/>', b'<w:altChunk r:id="html-body"/>')
        return members
    rewrite_parts(source, transform)
    result = extract_docx(source, tmp_path / 'out')
    assert not result.media_paths and not result.tables
    assert result.diagnostics and '본문 이미지 참조' in result.diagnostics[0]
    if kind == 'altChunk':
        assert 'altChunk' in result.diagnostics[1]
    with pytest.raises(ValueError, match='내장 미디어'):
        build_pptx(collect_project(tmp_path), tmp_path / 'report.pptx')
