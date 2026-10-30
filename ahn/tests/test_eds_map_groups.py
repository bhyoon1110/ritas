from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import sys

from PIL import Image
from pptx import Presentation
import pytest

from ahn.analysis import collect_project
from ahn.docx_extract import DocxImage, extract_docx
from ahn.eds_map import _read_electron_label, split_map_groups
from ahn.ppt_report import build_pptx
from ahn.tests.test_ppt_report import _body_slides


@pytest.fixture(autouse=True)
def no_optional_ocr(monkeypatch):
    monkeypatch.setattr('ahn.eds_map.shutil.which', lambda name: None)


def write_map_docx(path, counts=(10, 10), *, headings=False, same_pixels=False, strict=False):
    """Instrument layout: large electron image, then a paragraph of small maps."""
    media, paragraphs, relationships, colors = {}, [], [], []
    for group, count in enumerate(counts, 1):
        if headings:
            paragraphs.append(f'<w:p><w:r><w:t>Electron </w:t></w:r><w:r><w:t>Image {group}</w:t></w:r></w:p>')
        drawings = []
        for index in range(count + 1):
            rid = f'image{group}_{index}'
            color = (group * 45, 35 + index * 10, 60)
            colors.append(color)
            output = BytesIO()
            size = (900, 560) if index == 0 else (280, 250)
            Image.new('RGB', (900, 560) if same_pixels else size, color).save(output, format='PNG')
            media[f'word/media/{rid}.png'] = output.getvalue()
            relationships.append(f'<Relationship Id="{rid}" Target="media/{rid}.png"/>')
            drawing = (f'<w:r><w:drawing><wp:inline><wp:extent cx="{size[0]*1000}" cy="{size[1]*1000}"/>'
                       f'<a:blip r:embed="{rid}"/></wp:inline></w:drawing></w:r>')
            if index == 0:
                paragraphs.append('<w:p>' + drawing + '</w:p>')
            else:
                drawings.append(drawing)
        paragraphs.append('<w:p>' + ''.join(drawings) + '</w:p>')
        paragraphs.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
           '<w:body>' + ''.join(paragraphs) + '</w:body></w:document>')
    if strict:
        for old, new in (
            ('http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'http://purl.oclc.org/ooxml/wordprocessingml/main'),
            ('http://schemas.openxmlformats.org/drawingml/2006/main', 'http://purl.oclc.org/ooxml/drawingml/main'),
            ('http://schemas.openxmlformats.org/officeDocument/2006/relationships', 'http://purl.oclc.org/ooxml/officeDocument/relationships'),
            ('http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing', 'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing'),
        ):
            xml = xml.replace(old, new)
    with ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', xml)
        archive.writestr('word/_rels/document.xml.rels',
                        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                        + ''.join(relationships) + '</Relationships>')
        for name, data in media.items():
            archive.writestr(name, data)
    return colors


@pytest.mark.parametrize('headings,same_pixels,strict', [(False, False, False), (False, True, False), (True, True, True)])
def test_two_datasets_keep_each_anchor_and_all_maps_across_continuation_pages(tmp_path, headings, same_pixels, strict):
    source = tmp_path / 'sample MAP.docx'
    colors = write_map_docx(source, headings=headings, same_pixels=same_pixels, strict=strict)
    original = source.read_bytes()
    extracted = extract_docx(source, tmp_path / 'extracted')
    groups = split_map_groups(extracted.image_items)
    assert [group.label for group in groups] == ['Electron Image1', 'Electron Image2']
    assert [len(group.images) for group in groups] == [11, 11]
    output = build_pptx(collect_project(tmp_path), tmp_path / 'result.pptx')
    presentation = Presentation(output)
    slides = _body_slides(presentation)
    assert len(slides) == 4
    expected = [colors[:7], [colors[0], *colors[7:11]], colors[11:18], [colors[11], *colors[18:]]]
    for index, (slide, expected_colors) in enumerate(zip(slides, expected)):
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        actual_colors = [Image.open(BytesIO(shape.image.blob)).convert('RGB').getpixel((10, 10)) for shape in pictures]
        assert len(actual_colors) == len(expected_colors)
        for actual, expected in zip(actual_colors, expected_colors):
            assert actual == pytest.approx(expected, abs=2)  # PPT image cache uses JPEG.
        assert any(shape.has_text_frame and f'Electron Image{index//2+1}' in shape.text for shape in slide.shapes)
        title_bottom = max(shape.top + shape.height for shape in slide.shapes
                           if shape.has_text_frame and 'STEM EDS 분석결과' in shape.text)
        for shape in pictures:
            assert 0 <= shape.left <= presentation.slide_width - shape.width
            assert title_bottom < shape.top <= presentation.slide_height - shape.height
    assert source.read_bytes() == original


def test_three_unequal_groups_and_pixel_fallback_keep_order(tmp_path):
    source = tmp_path / 'three MAP.docx'
    write_map_docx(source, counts=(2, 7, 1))
    extracted = extract_docx(source, tmp_path / 'images')
    groups = split_map_groups([DocxImage(path) for path in extracted.media_paths])
    assert [len(g.images) for g in groups] == [3, 8, 2]
    assert [path for group in groups for path in group.images] == extracted.media_paths


def test_explicit_labels_split_same_size_images_and_preserve_numbers(tmp_path):
    image = tmp_path / 'same.png'
    Image.new('RGB', (100, 100), 'gray').save(image)
    items = [DocxImage(image, label) for label in ('Electron Image7', '', '', 'Electron Image22', '')]
    groups = split_map_groups(items)
    assert [(g.label, len(g.images)) for g in groups] == [('Electron Image7', 3), ('Electron Image22', 2)]


def test_equal_size_maps_are_not_split_by_count_or_moderate_size_variation(tmp_path):
    items = []
    for index in range(16):
        path = tmp_path / f'{index}.png'
        Image.new('RGB', (300 if index == 8 else 280, 250), 'gray').save(path)
        items.append(DocxImage(path))
    groups = split_map_groups(items)
    assert len(groups) == 1 and groups[0].label == '' and len(groups[0].images) == 16
    assert split_map_groups([]) == []


def test_reused_media_is_kept_for_each_dataset_but_alternate_content_is_once(tmp_path):
    source = tmp_path / 'repeat MAP.docx'
    write_map_docx(source, counts=(1,))
    with ZipFile(source) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    xml = parts['word/document.xml'].decode()
    xml = xml.replace('<w:body>', '<w:body xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                      'xmlns:v="urn:schemas-microsoft-com:vml">')
    xml = xml.replace('<a:blip r:embed="image1_0"/>', '<mc:AlternateContent><mc:Choice Requires="a">'
                      '<a:blip r:embed="image1_0"/></mc:Choice><mc:Fallback><v:imagedata r:id="image1_0"/>'
                      '</mc:Fallback></mc:AlternateContent>')
    xml = xml.replace('</w:body>', '<w:p><w:r><w:t>Electron Image 2</w:t></w:r></w:p>'
                      '<w:p><w:r><w:drawing><a:blip r:embed="image1_0"/></w:drawing></w:r></w:p>'
                      '<w:p><w:r><w:drawing><a:blip r:embed="image1_1"/></w:drawing></w:r></w:p></w:body>')
    parts['word/document.xml'] = xml.encode()
    with ZipFile(source, 'w') as archive:
        for name, data in parts.items(): archive.writestr(name, data)
    result = extract_docx(source, tmp_path / 'out')
    assert len(result.media_paths) == 2
    assert [item.path for item in result.image_items] == result.media_paths * 2
    assert [len(group.images) for group in split_map_groups(result.image_items)] == [2, 2]


def test_title_ocr_composites_transparency_and_is_bounded(tmp_path, monkeypatch):
    source = tmp_path / 'anchor.png'
    Image.new('RGBA', (900, 560), (0, 0, 0, 0)).save(source)
    calls = []
    def read(image, **kwargs):
        calls.append(kwargs)
        assert image.getpixel((0, 0)) == 255
        assert image.width <= 1800 and image.height < 140
        return 'Electron Image 32\nMap Data 1'
    monkeypatch.setattr('ahn.eds_map.shutil.which', lambda name: '/mock/tesseract')
    monkeypatch.setitem(sys.modules, 'pytesseract', SimpleNamespace(image_to_string=read))
    assert _read_electron_label(source) == 'Electron Image32'
    assert calls == [{'lang': 'eng', 'config': '--psm 6', 'timeout': 2}]
    def timeout(*args, **kwargs): raise RuntimeError('timeout')
    monkeypatch.setitem(sys.modules, 'pytesseract', SimpleNamespace(image_to_string=timeout))
    assert _read_electron_label(source) == ''
